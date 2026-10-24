"""Build site-migration-checklist workbook (.xlsx) and checklist.md from checks.csv.

    pip install openpyxl
    python build.py

Part of https://github.com/lazarinastoy/mlforseo-code (MIT). Lazarina Stoy, MLforSEO.
"""
# ---------------------------------------------------------------------------
# Checklist workbook engine (shared by the MLforSEO checklist folders).
# Builds a Google-Sheets-ready .xlsx with openpyxl and a plain checklist.md
# from checks.csv. Only uses functions that work in both Google Sheets and
# Excel (COUNTIF, COUNTIFS, COUNTA, IF, IFERROR, ROUND, REPT, HYPERLINK, MID,
# FIND, LEN, SUBSTITUTE, RIGHT, AND, OR).
# ---------------------------------------------------------------------------
import csv
import math
import pathlib

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = pathlib.Path(__file__).resolve().parent

# MLforSEO palette, used sparingly
C = {
    "header": "251046", "ink": "1A0B32", "accent": "8448EF", "band": "F4F0FD",
    "high": "FFDE59", "white": "FFFFFF", "muted": "6B6480", "line": "DCD3F0",
    "done_fill": "DDF3E4", "done_font": "1D6B3A", "prog_fill": "E6DCFB",
    "prog_font": "4B2A99", "na_fill": "EEEEEE", "na_font": "8A8A8A",
    "warn_fill": "FDE2E2", "warn_font": "9B1C1C",
}
FONT = "Arial"  # Google Sheets' default; renders the same in Excel
STATUSES = ["Not started", "In progress", "Done", "N/A"]
PRIORITIES = ["High", "Medium", "Low"]
EVIDENCE = [
    ("Google-documented", "Stated in Google Search Central or Google Help documentation."),
    ("Platform-documented", "Stated in another platform's own documentation (OpenAI, Anthropic, Perplexity, Microsoft Bing, Apple, Meta, Wikidata / Wikimedia, schema.org)."),
    ("Research-backed", "Supported by published research or a large public dataset. Results come from specific test conditions and may not carry over to every site."),
    ("MLforSEO practice", "What we teach in MLforSEO courses and use in audits, based on how retrieval and entity systems work. Sensible and low-risk, but not confirmed by the platforms."),
    ("Unconfirmed (optional)", "A cheap extra that no major platform has confirmed using. Do it last, if at all."),
]

THIN = Side(style="thin", color=C["line"])
BORDER = Border(bottom=THIN)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
CENTER_TOP = Alignment(horizontal="center", vertical="top", wrap_text=True)


def font(size=10, bold=False, color=None, italic=False, underline=None):
    return Font(name=FONT, size=size, bold=bold, color=color or C["ink"], italic=italic, underline=underline)


def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


def q(text):
    """Escape a string for use inside a spreadsheet formula."""
    return str(text).replace('"', '""')


def link_formula(url, label):
    if not url:
        return ""
    return f'=HYPERLINK("{q(url)}","{q(label or url)}")'


def est_lines(text, width_chars):
    if not text:
        return 1
    lines = 0
    for para in str(text).split("\n"):
        lines += max(1, math.ceil(len(para) * 1.08 / max(width_chars, 1)))
    return lines


def read_checks(path=None):
    path = path or HERE / "checks.csv"
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


# --------------------------------------------------------------------------
# Generic table helpers
# --------------------------------------------------------------------------
def header_row(ws, headers, row=1, widths=None, height=30):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = font(10, True, C["white"])
        c.fill = fill(C["header"])
        c.alignment = Alignment(wrap_text=True, vertical="center")
        if widths:
            ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.row_dimensions[row].height = height


def list_validation(ws, values, rng, prompt=None):
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True,
                        showDropDown=False)  # showDropDown=False means the in-cell arrow IS shown
    dv.error = "Pick a value from the list."
    dv.errorTitle = "Not in list"
    if prompt:
        dv.prompt = prompt
    ws.add_data_validation(dv)
    dv.add(rng)
    return dv


def status_cf(ws, rng):
    rules = [
        ("Done", C["done_fill"], C["done_font"]),
        ("In progress", C["prog_fill"], C["prog_font"]),
        ("N/A", C["na_fill"], C["na_font"]),
        ("Not started", None, C["muted"]),
    ]
    for value, f_, col in rules:
        kw = {"font": Font(name=FONT, color=col, bold=value != "Not started")}
        if f_:
            kw["fill"] = fill(f_)
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{value}"'], stopIfTrue=True, **kw))


def priority_cf(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"High"'], stopIfTrue=True,
                                                  fill=fill(C["high"]), font=Font(name=FONT, bold=True, color=C["header"])))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Low"'], stopIfTrue=True,
                                                  font=Font(name=FONT, color=C["muted"])))


def banding_cf(ws, rng):
    # Added last so status / priority colours win. Survives sorting and filtering.
    ws.conditional_formatting.add(rng, FormulaRule(formula=["MOD(ROW(),2)=0"], fill=fill(C["band"])))


def print_setup(ws, landscape=True):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.gridLines = False


# --------------------------------------------------------------------------
# Checklist tab
# --------------------------------------------------------------------------
CHECK_COLS = [
    ("ID", 7), (None, 20), ("Check", 44), ("Why it matters", 46), ("How to check (tool / where)", 46),
    ("Priority", 10), ("Evidence", 17), ("Status", 13), ("Owner", 13), ("Notes", 28),
    ("Learn more", 30), ("Official source", 30),
]


def build_checklist(wb, checks, title="Checklist", area_label="Area", spare_rows=15, tab_color=None):
    ws = wb.create_sheet(title)
    cols = [(area_label if h is None else h, w) for h, w in CHECK_COLS]
    header_row(ws, [h for h, _ in cols], widths=[w for _, w in cols])
    r = 1
    for ck in checks:
        r += 1
        values = [ck["id"], ck["area"], ck["check"], ck["why"], ck["how"], ck["priority"], ck["evidence"],
                  "Not started", "", "", link_formula(ck.get("learn_url"), ck.get("learn_label")),
                  link_formula(ck.get("source_url"), ck.get("source_label"))]
        for ci, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=ci, value=v)
            cell.font = font(10, bold=(ci == 3))
            cell.alignment = CENTER_TOP if ci in (1, 6, 8) else WRAP_TOP
            cell.border = BORDER
            if ci in (11, 12) and v:
                cell.font = font(10, color=C["accent"], underline="single")
        lines = max(est_lines(ck["check"], 44), est_lines(ck["why"], 46), est_lines(ck["how"], 46))
        ws.row_dimensions[r].height = min(15 * lines + 6, 300)
    last_check = r
    for _ in range(spare_rows):  # blank rows for the user's own checks
        r += 1
        for ci in range(1, len(cols) + 1):
            cell = ws.cell(row=r, column=ci, value=None)
            cell.font = font(10)
            cell.alignment = WRAP_TOP
            cell.border = BORDER
    last = r
    ncol = get_column_letter(len(cols))
    list_validation(ws, STATUSES, f"H2:H{last}")
    list_validation(ws, PRIORITIES, f"F2:F{last}")
    list_validation(ws, [e for e, _ in EVIDENCE], f"G2:G{last}")
    status_cf(ws, f"H2:H{last}")
    priority_cf(ws, f"F2:F{last}")
    banding_cf(ws, f"A2:{ncol}{last}")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{ncol}{last}"
    if tab_color:
        ws.sheet_properties.tabColor = tab_color
    print_setup(ws)
    return ws, last_check, last


# --------------------------------------------------------------------------
# Start here tab
# --------------------------------------------------------------------------
def build_start(wb, blocks, tab_color=C["accent"]):
    """blocks: list of tuples. Types: title, subtitle, h, p, steps, kv (list of (k, v)),
    status_legend, priority_legend, evidence_legend, links (list of (label, url)), spacer."""
    ws = wb.active
    ws.title = "Start here"
    ws.sheet_properties.tabColor = tab_color
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 92
    r = 1

    def para(text, size=10, bold=False, color=None, italic=False, height=None):
        nonlocal r
        r += 1
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        c = ws.cell(row=r, column=2, value=text)
        c.font = font(size, bold, color, italic)
        c.alignment = WRAP_TOP
        ws.row_dimensions[r].height = height or (max(1, est_lines(text, 112)) * (size + 5) + 4)

    def kv(k, v, kfill=None, kfont=None, link=None):
        nonlocal r
        r += 1
        a = ws.cell(row=r, column=2, value=k)
        a.font = kfont or font(10, True)
        a.alignment = WRAP_TOP
        if kfill:
            a.fill = fill(kfill)
        b = ws.cell(row=r, column=3, value=link_formula(link, v) if link else v)
        b.font = font(10, color=C["accent"], underline="single") if link else font(10)
        b.alignment = WRAP_TOP
        ws.row_dimensions[r].height = max(est_lines(k, 26), est_lines(v, 90)) * 15 + 4

    for blk in blocks:
        kind = blk[0]
        if kind == "title":
            para(blk[1], 18, True, C["header"], height=30)
        elif kind == "subtitle":
            para(blk[1], 11, False, C["muted"])
        elif kind == "h":
            r += 1
            para(blk[1], 12, True, C["accent"], height=22)
        elif kind == "p":
            para(blk[1])
        elif kind == "note":
            para(blk[1], 9, False, C["muted"], italic=True)
        elif kind == "steps":
            for i, s in enumerate(blk[1], start=1):
                kv(f"Step {i}", s)
        elif kind == "kv":
            for k, v in blk[1]:
                kv(k, v)
        elif kind == "status_legend":
            kv("Not started", "Default for every check.", kfont=Font(name=FONT, color=C["muted"]))
            kv("In progress", "Someone is working on it.", kfill=C["prog_fill"], kfont=Font(name=FONT, bold=True, color=C["prog_font"]))
            kv("Done", "Checked, and fixed if it needed fixing.", kfill=C["done_fill"], kfont=Font(name=FONT, bold=True, color=C["done_font"]))
            kv("N/A", "Does not apply to this site. N/A rows are left out of the % complete.", kfill=C["na_fill"], kfont=Font(name=FONT, color=C["na_font"]))
        elif kind == "priority_legend":
            kv("High", blk[1][0], kfill=C["high"], kfont=Font(name=FONT, bold=True, color=C["header"]))
            kv("Medium", blk[1][1])
            kv("Low", blk[1][2], kfont=Font(name=FONT, color=C["muted"]))
        elif kind == "evidence_legend":
            for k, v in EVIDENCE:
                kv(k, v)
        elif kind == "links":
            for label, url in blk[1]:
                kv(label[0], label[1], link=url)
        elif kind == "spacer":
            r += 1
    print_setup(ws, landscape=False)
    return ws


# --------------------------------------------------------------------------
# Progress tab (COUNTIF / COUNTIFS only, so it works in Sheets and Excel)
# --------------------------------------------------------------------------
def build_progress(wb, sheet, areas, area_label="Area", last_row=500, title="Progress"):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = C["high"]
    s = f"'{sheet}'"
    A, P, S, ID = f"{s}!$B$2:$B${last_row}", f"{s}!$F$2:$F${last_row}", f"{s}!$H$2:$H${last_row}", f"{s}!$C$2:$C${last_row}"
    widths = [3, 30, 10, 10, 13, 13, 10, 12, 28, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws["B2"] = "Progress"
    ws["B2"].font = font(18, True, C["header"])
    ws["B3"] = "Updates automatically as you change the Status column. Formulas use COUNTIF / COUNTIFS, so they work in Google Sheets and Excel."
    ws["B3"].font = font(9, color=C["muted"], italic=True)

    # Overall
    labels = [("Checks", f"=COUNTA({ID})"), ("Done", f'=COUNTIF({S},"Done")'),
              ("In progress", f'=COUNTIF({S},"In progress")'), ("Not started", f'=COUNTIF({S},"Not started")'),
              ("N/A", f'=COUNTIF({S},"N/A")')]
    r = 5
    for k, f_ in labels:
        ws.cell(row=r, column=2, value=k).font = font(10, True)
        c = ws.cell(row=r, column=3, value=f_)
        c.font = font(10)
        c.alignment = Alignment(horizontal="right")
        r += 1
    ws.cell(row=r, column=2, value="% complete (Done / applicable)").font = font(10, True)
    pc = ws.cell(row=r, column=3, value="=IFERROR(C6/(C5-C9),0)")
    pc.number_format = "0%"
    pc.font = font(12, True, C["accent"])
    bar = ws.cell(row=r, column=4, value=f'=REPT("█",ROUND(C{r}*20,0))&REPT("░",20-ROUND(C{r}*20,0))')
    bar.font = Font(name=FONT, size=11, color=C["accent"])
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=8)
    pct_row = r
    r += 1
    ws.cell(row=r, column=2, value="Open High-priority checks").font = font(10, True)
    c = ws.cell(row=r, column=3, value=f'=COUNTIFS({P},"High",{S},"<>Done",{S},"<>N/A")')
    c.font = font(12, True, C["header"])
    c.fill = fill(C["high"])

    # By area
    r += 2
    hdr = [area_label, "Checks", "Done", "In progress", "Not started", "N/A", "% done", "Open High", ]
    for i, h in enumerate(hdr, start=2):
        c = ws.cell(row=r, column=i, value=h)
        c.font = font(10, True, C["white"])
        c.fill = fill(C["header"])
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[r].height = 28
    first = r + 1
    for a in areas:
        r += 1
        ws.cell(row=r, column=2, value=a).font = font(10, True)
        ref = f"$B{r}"
        fs = [f"=COUNTIF({A},{ref})",
              f'=COUNTIFS({A},{ref},{S},"Done")', f'=COUNTIFS({A},{ref},{S},"In progress")',
              f'=COUNTIFS({A},{ref},{S},"Not started")', f'=COUNTIFS({A},{ref},{S},"N/A")',
              f"=IFERROR(D{r}/(C{r}-G{r}),0)",
              f'=COUNTIFS({A},{ref},{P},"High",{S},"<>Done",{S},"<>N/A")']
        for i, f_ in enumerate(fs, start=3):
            c = ws.cell(row=r, column=i, value=f_)
            c.font = font(10)
            c.alignment = Alignment(horizontal="right")
            if i == 8:
                c.number_format = "0%"
        bar = ws.cell(row=r, column=10, value=f'=REPT("█",ROUND(H{r}*10,0))&REPT("░",10-ROUND(H{r}*10,0))')
        bar.font = Font(name=FONT, size=10, color=C["accent"])
    ws.column_dimensions["J"].width = 16
    ws.conditional_formatting.add(f"B{first}:J{r}", FormulaRule(formula=["MOD(ROW(),2)=0"], fill=fill(C["band"])))
    ws.conditional_formatting.add(f"I{first}:I{r}", CellIsRule(operator="greaterThan", formula=["0"],
                                                               fill=fill(C["high"]), font=Font(name=FONT, bold=True, color=C["header"])))

    # By priority
    r += 2
    for i, h in enumerate(["Priority", "Checks", "Done", "Open", "% done"], start=2):
        c = ws.cell(row=r, column=i, value=h)
        c.font = font(10, True, C["white"])
        c.fill = fill(C["header"])
    for p in PRIORITIES:
        r += 1
        ws.cell(row=r, column=2, value=p).font = font(10, True)
        ref = f"$B{r}"
        fs = [f"=COUNTIF({P},{ref})", f'=COUNTIFS({P},{ref},{S},"Done")',
              f'=COUNTIFS({P},{ref},{S},"<>Done",{S},"<>N/A")',
              f'=IFERROR(D{r}/(C{r}-COUNTIFS({P},{ref},{S},"N/A")),0)']
        for i, f_ in enumerate(fs, start=3):
            c = ws.cell(row=r, column=i, value=f_)
            c.font = font(10)
            c.alignment = Alignment(horizontal="right")
            if i == 6:
                c.number_format = "0%"
    priority_cf(ws, f"B{r-2}:B{r}")
    print_setup(ws, landscape=False)
    return ws, r


# --------------------------------------------------------------------------
# Simple log / tracker tabs
# --------------------------------------------------------------------------
def build_log(wb, title, columns, rows=200, validations=None, intro=None, example_rows=None, tab_color=None,
              formulas=None, wrap_cols=None):
    """columns: list of (header, width). validations: {col_letter: [values]}.
    formulas: {col_letter: template with {r}} filled for every row."""
    ws = wb.create_sheet(title)
    start = 1
    if intro:
        ws.cell(row=1, column=1, value=intro).font = font(9, color=C["muted"], italic=True)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
        ws.row_dimensions[1].height = max(1, est_lines(intro, sum(w for _, w in columns))) * 14 + 6
        ws.cell(row=1, column=1).alignment = WRAP_TOP
        start = 2
    header_row(ws, [h for h, _ in columns], row=start, widths=[w for _, w in columns])
    first = start + 1
    last = start + rows
    example_rows = example_rows or []
    for i in range(rows):
        r = first + i
        ex = example_rows[i] if i < len(example_rows) else None
        for ci in range(1, len(columns) + 1):
            col = get_column_letter(ci)
            v = None
            if formulas and col in formulas:
                v = formulas[col].format(r=r)
            elif ex is not None and ci - 1 < len(ex):
                v = ex[ci - 1]
            cell = ws.cell(row=r, column=ci, value=v if v != "" else None)
            cell.font = font(10, italic=ex is not None and not (formulas and col in formulas),
                             color=C["muted"] if ex is not None else None)
            cell.alignment = WRAP_TOP if (wrap_cols is None or col in wrap_cols) else Alignment(vertical="top")
            cell.border = BORDER
    for col, values in (validations or {}).items():
        list_validation(ws, values, f"{col}{first}:{col}{last}")
    ncol = get_column_letter(len(columns))
    ws.freeze_panes = ws.cell(row=first, column=2)
    ws.auto_filter.ref = f"A{start}:{ncol}{last}"
    if tab_color:
        ws.sheet_properties.tabColor = tab_color
    print_setup(ws)
    return ws, first, last


# --------------------------------------------------------------------------
# Markdown version
# --------------------------------------------------------------------------
def md_link(label, url):
    return f"[{label}]({url})" if url else ""


def write_markdown(path, title, intro_md, checks, area_order, outro_md, area_intros=None):
    out = [f"# {title}", "", intro_md.strip(), ""]
    for area in area_order:
        out.append(f"## {area}")
        out.append("")
        if area_intros and area in area_intros:
            out += [area_intros[area], ""]
        for ck in [c for c in checks if c["area"] == area]:
            out.append(f"- [ ] **{ck['id']} · {ck['check']}**  ")
            out.append(f"  *Priority: {ck['priority']} · Evidence: {ck['evidence']}*  ")
            out.append(f"  **Why it matters:** {ck['why']}  ")
            out.append(f"  **How to check:** {ck['how']}  ")
            links = [x for x in (md_link(ck.get('learn_label'), ck.get('learn_url')),
                                 md_link(ck.get('source_label'), ck.get('source_url'))) if x]
            if links:
                out.append(f"  **Learn more:** {' · '.join(links)}")
            out.append("")
    out += [outro_md.strip(), ""]
    pathlib.Path(path).write_text("\n".join(out), encoding="utf-8")


# ===========================================================================
# SEO site migration checklist + redirect mapping workbook
# Run:  python build.py   (needs openpyxl)  ->  site-migration-checklist.xlsx + checklist.md
# ===========================================================================
TITLE = "SEO Site Migration Checklist + Redirect Mapping Workbook"
XLSX = "site-migration-checklist.xlsx"
AREAS = ['1. Scope & plan', '2. Benchmark', '3. Redirect mapping', '4. Pre-launch (staging)', '5. Launch day', '6. First weeks', '7. Months 1 to 12']
DECISIONS = ["Accept suggestion", "Override", "410 Gone", "404", "No change (same URL)"]
QA = ["To review", "Approved", "Live and verified", "Problem"]
REDIRECT_ROWS = 2000

# Column layout of the Redirect map tab (letter, header, width, kind)
RM = [
    ("A", "Old URL", 40, "in"), ("B", "Clicks (last 12 months)", 11, "in"), ("C", "Referring domains", 11, "in"),
    ("D", "Old URL priority", 11, "f"), ("E", "Matched on", 14, "in"), ("F", "Suggested new URL (from notebook)", 40, "in"),
    ("G", "Similarity (0 to 100)", 11, "in"), ("H", "Match type (from notebook)", 16, "in"), ("I", "Algorithm", 14, "in"),
    ("J", "Algorithms agreeing", 11, "in"), ("K", "Decision", 18, "in"), ("L", "Override URL", 34, "in"),
    ("M", "Final destination", 40, "f"), ("N", "Redirect type", 9, "f"), ("O", "Source path", 30, "f"),
    ("P", "Apache rule (check before use)", 44, "f"), ("Q", "Status code at test", 10, "in"), ("R", "Hops", 7, "in"),
    ("S", "Final URL returns 200?", 10, "in"), ("T", "Flags", 30, "f"), ("U", "QA status", 15, "in"),
    ("V", "Reviewer", 12, "in"), ("W", "Notes", 30, "in"),
]

START = [
    ("title", TITLE),
    ("subtitle", "43 checks across seven phases, plus a redirect map that takes the output of MLforSEO's fuzzy-matching notebooks and tracks QA for every URL. Free from MLforSEO, October 2026."),
    ("h", "What this is"),
    ("p", "A phase-by-phase migration checklist (scope, benchmark, redirect mapping, staging, launch day, first weeks, first year) built on Google's site move documentation, "
          "linked to a redirect mapping workbook. Draft your redirects with fuzzy matching, paste the results in, make a decision for each old URL, and the workbook builds the final destination, "
          "an Apache rule and warning flags (home page redirects, chains, self-redirects, low scores), then tracks QA through to 'Live and verified'."),
    ("p", "It covers domain changes, HTTP to HTTPS, URL structure changes and CMS replatforms. If your URLs aren't changing (a pure hosting move), follow Google's guide for moves without URL changes instead."),
    ("h", "How to use it"),
    ("steps", [
        "Make your own copy (File > Make a copy in Google Sheets).",
        "Work through the Checklist tab phase by phase. Set Status, Owner and Notes as you go; mark anything that doesn't apply as N/A.",
        "In phase 2, paste every old URL into column A of the Redirect map, with clicks and referring domains. Set your thresholds in the yellow cells at the top of that tab.",
        "In phase 3, run a fuzzy-matching tool (see 'Which tool?' below) and paste its suggestions into columns E to J. Delete the example rows first.",
        "Review each row and pick a Decision. Accepted rows use the suggestion; Override rows use the URL you type in column L. Columns M to P and T fill in automatically.",
        "Before and after launch, crawl the old URLs in list mode and fill in Status code, Hops and 'Final URL returns 200?'. Move rows through QA status until they're 'Live and verified'.",
        "Watch the Progress tab for open High-priority checks, undecided URLs and flags.",
    ]),
    ("h", "Which fuzzy-matching tool?"),
    ("kv", [
        ("Under ~300 URLs", "Fuzzy Matching for SEO Google Sheets template (URL Redirection Mapping tab): www.mlforseo.com/resource/fuzzy-matching-for-seo-google-sheets-template/"),
        ("Old site vs new site", "Triple fuzzy matching redirect notebook: matches pre- and post-migration crawls on Address, Title 1, H1-1, H2-1, H2-2 and Meta Description 1 with three algorithms."),
        ("404s to live pages", "Mapping 404 URLs to live pages notebook: four algorithms on the URL path, outputs .htaccess rules (scores are 0 to 1, so multiply by 100 for column G)."),
        ("Python, free on GitHub", "fuzzy-matching-for-seo folder in lazarinastoy/mlforseo-code: redirect section outputs old_url, match, new_url, similarity, action."),
    ]),
    ("h", "Pasting notebook output into the Redirect map"),
    ("kv", [
        ("Old URL (A)", "Pre-Migration value (triple notebook, Address rows) · 404 URL (404 notebook) · old_url (GitHub notebook) · old URL column (Sheets template)"),
        ("Matched on (E)", "The column that generated the match: Address, Title 1, H1-1, H2-1, H2-2 or Meta Description 1. For title or heading matches, look up the URL that belongs to the matched value in your crawl before pasting."),
        ("Suggested new URL (F)", "Post-Migration match · live URL · new_url · most similar new URL"),
        ("Similarity (G)", "Similarity score on a 0 to 100 scale (multiply 0 to 1 scores by 100)"),
        ("Match type (H)", "Exact Match / Partial Match / No Suitable Match Found, or direct match / fuzzy match"),
        ("Algorithm (I), agreeing (J)", "PolyFuzz (TF-IDF), RapidFuzz, TheFuzz / FuzzyWuzzy, Levenshtein. Put the number of algorithms that suggested the same URL in J: agreement is a good sign."),
    ]),
    ("h", "Status"),
    ("status_legend",),
    ("h", "Priority"),
    ("priority_legend", ["Skipping it risks lost traffic, broken indexing or lost link value.",
                         "Reduces risk and speeds up recovery.",
                         "Tidy-up and documentation."]),
    ("h", "Evidence levels"),
    ("evidence_legend",),
    ("h", "About the Apache rule column"),
    ("note", "Column P writes mod_alias rules such as 'Redirect 301 /old-page/ https://www.example.com/new-page/'. Redirect matches by prefix (it also catches /old-page/anything) and ignores query strings. "
             "Check the rules with your developer, order specific rules before general ones, or use RedirectMatch, nginx, CDN or CMS redirects instead. Columns M and O (destination and source path) work for any platform."),
    ("h", "Credits"),
    ("links", [
        (("Created by", "Lazarina Stoy · MLforSEO"), "https://www.mlforseo.com/experts/lazarina-stoy/"),
        (("Read first", "404 and redirect mapping with fuzzy matching"), "https://www.mlforseo.com/machine-learning-tutorials/fuzzy-matching/404-and-redirect-mapping-with-fuzzy-matching-in-google-colab-python/"),
        (("Official guide", "Google: how to move a site with URL changes"), "https://developers.google.com/search/docs/crawling-indexing/site-move-with-url-changes"),
        (("Go deeper", "Introduction to Machine Learning for SEO (fuzzy matching module)"), "https://www.mlforseo.com/courses/introduction-to-machine-learning-for-seo/"),
        (("Source files", "GitHub: lazarinastoy/mlforseo-code"), "https://github.com/lazarinastoy/mlforseo-code/tree/main/site-migration-checklist"),
    ]),
    ("note", "Version 1.0, October 2026. Free to use and adapt (MIT licence); please keep the credit."),
]

EXAMPLES = [
    ["https://old.example.com/blog/how-to-choose-running-shoes/", 420, 12, None, "Address", "https://www.example.com/guides/choosing-running-shoes/", 82, "Partial Match", "RapidFuzz", 3, "Accept suggestion", "", None, None, None, None, 301, 1, "Yes", None, "Live and verified", "LS", "EXAMPLE ROW: delete."],
    ["https://old.example.com/products/trail-runner-2/", 95, 3, None, "Title 1", "https://www.example.com/shop/trail-runner-3/", 71, "Partial Match", "PolyFuzz (TF-IDF)", 1, "Override", "https://www.example.com/shop/trail-runners/", None, None, None, None, None, None, None, None, "Approved", "LS", "EXAMPLE ROW: delete. Model discontinued, category is the closest match."],
    ["https://old.example.com/2019-summer-sale/", 0, 0, None, "Address", "https://www.example.com/", 38, "No Suitable Match Found", "RapidFuzz", 1, "410 Gone", "", None, None, None, None, None, None, None, None, "Approved", "LS", "EXAMPLE ROW: delete. Expired promotion."],
    ["https://old.example.com/about-us/", 60, 25, None, "H1-1", "https://www.example.com/about/", 100, "Exact Match", "TheFuzz", 3, "", "", None, None, None, None, None, None, None, None, "To review", "", "EXAMPLE ROW: delete."],
]


def rm_formulas(r):
    thr = "$C$2"   # clicks threshold
    rdt = "$E$2"   # referring domains threshold
    sim = "$G$2"   # similarity threshold
    slashes = f'(LEN(M{r})-LEN(SUBSTITUTE(M{r},"/","")))'
    return {
        "D": f'=IF(A{r}="","",IF(OR(N(B{r})>={thr},N(C{r})>={rdt}),"High",IF(OR(N(B{r})>0,N(C{r})>0),"Medium","Low")))',
        "M": f'=IF(K{r}="Override",L{r},IF(K{r}="Accept suggestion",F{r},IF(K{r}="No change (same URL)",A{r},"")))',
        "N": f'=IF(A{r}="","",IF(K{r}="410 Gone","410",IF(K{r}="404","404",IF(K{r}="No change (same URL)","none",IF(M{r}="","","301")))))',
        "O": f'=IF(A{r}="","",IFERROR(MID(A{r},FIND("/",A{r},FIND("//",A{r})+2),2000),A{r}))',
        "P": f'=IF(OR(O{r}="",N{r}="",N{r}="none",N{r}="404"),"",IF(N{r}="410","Redirect gone "&O{r},"Redirect "&N{r}&" "&O{r}&" "&M{r}))',
        "T": (f'=IF(A{r}="","",TRIM('
              f'IF(K{r}="","Needs decision; ","")&'
              f'IF(AND(K{r}="Accept suggestion",F{r}=""),"No suggestion to accept; ","")&'
              f'IF(AND(K{r}="Override",L{r}=""),"Override URL missing; ","")&'
              f'IF(AND(K{r}="Accept suggestion",N(G{r})<{sim}),"Below your threshold; ","")&'
              f'IF(AND(M{r}<>"",M{r}=A{r},K{r}<>"No change (same URL)"),"Redirects to itself; ","")&'
              f'IF(AND(M{r}<>"",OR(M{r}="/",AND(LEFT(M{r},4)="http",OR({slashes}=2,AND({slashes}=3,RIGHT(M{r},1)="/"))))),"Home page redirect; ","")&'
              f'IF(N(R{r})>1,"Chain; ","")&'
              f'IF(S{r}="No","Destination not 200; ","")))'),
    }


def build_redirect_map(wb):
    ws = wb.create_sheet("Redirect map")
    ws.sheet_properties.tabColor = C["accent"]
    ncol = len(RM)
    intro = ("Paste old URLs in column A and fuzzy-matching suggestions in E to J (see Start here for the column mapping). Pick a Decision in K. "
             "Grey-headed columns (D, M, N, O, P, T) are formulas: don't type over them. Delete the four example rows before you start.")
    c = ws.cell(row=1, column=1, value=intro)
    c.font = font(9, color=C["muted"], italic=True)
    c.alignment = WRAP_TOP
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    ws.row_dimensions[1].height = 30
    # Settings row
    settings = [("A2", "Settings (edit the yellow cells):"), ("B2", "High priority if clicks ≥"), ("C2", 100),
                ("D2", "or referring domains ≥"), ("E2", 5), ("F2", "Flag accepted matches with similarity below"), ("G2", 70)]
    for ref, v in settings:
        ws[ref] = v
        ws[ref].font = font(10, bold=isinstance(v, str))
        ws[ref].alignment = Alignment(wrap_text=True, vertical="center")
        if not isinstance(v, str):
            ws[ref].fill = fill(C["high"])
            ws[ref].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 30
    hr = 4
    for i, (col, h, w, kind) in enumerate(RM, start=1):
        cell = ws.cell(row=hr, column=i, value=h)
        cell.font = font(10, True, C["white"])
        cell.fill = fill(C["header"] if kind == "in" else C["muted"])
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[col].width = w
    ws.row_dimensions[hr].height = 44
    first, last = hr + 1, hr + REDIRECT_ROWS
    for i in range(REDIRECT_ROWS):
        r = first + i
        ex = EXAMPLES[i] if i < len(EXAMPLES) else None
        f = rm_formulas(r)
        for ci, (col, _, _, kind) in enumerate(RM, start=1):
            v = f[col] if col in f else (ex[ci - 1] if ex is not None else None)
            cell = ws.cell(row=r, column=ci, value=v if v != "" else None)
            cell.font = font(10, italic=ex is not None and kind == "in", color=C["muted"] if (ex is not None and kind == "in") else None)
            cell.alignment = Alignment(vertical="top", wrap_text=col in ("T", "W"))
    list_validation(ws, DECISIONS, f"K{first}:K{last}")
    list_validation(ws, QA, f"U{first}:U{last}")
    list_validation(ws, ["Yes", "No"], f"S{first}:S{last}")
    list_validation(ws, ["Address", "Title 1", "H1-1", "H2-1", "H2-2", "Meta Description 1", "URL path"], f"E{first}:E{last}")
    rng = lambda col: f"{col}{first}:{col}{last}"
    warn = dict(fill=fill(C["warn_fill"]), font=Font(name=FONT, color=C["warn_font"], bold=True))
    ws.conditional_formatting.add(rng("T"), FormulaRule(formula=[f'LEN(T{first})>0'], stopIfTrue=True, **warn))
    ws.conditional_formatting.add(rng("D"), CellIsRule(operator="equal", formula=['"High"'], stopIfTrue=True, fill=fill(C["high"]), font=Font(name=FONT, bold=True, color=C["header"])))
    ws.conditional_formatting.add(rng("G"), FormulaRule(formula=[f'AND(G{first}<>"",G{first}<$G$2)'], stopIfTrue=True, font=Font(name=FONT, color=C["warn_font"], bold=True)))
    ws.conditional_formatting.add(rng("U"), CellIsRule(operator="equal", formula=['"Live and verified"'], stopIfTrue=True, fill=fill(C["done_fill"]), font=Font(name=FONT, color=C["done_font"], bold=True)))
    ws.conditional_formatting.add(rng("U"), CellIsRule(operator="equal", formula=['"Approved"'], stopIfTrue=True, fill=fill(C["prog_fill"]), font=Font(name=FONT, color=C["prog_font"], bold=True)))
    ws.conditional_formatting.add(rng("U"), CellIsRule(operator="equal", formula=['"Problem"'], stopIfTrue=True, **warn))
    ws.conditional_formatting.add(rng("S"), CellIsRule(operator="equal", formula=['"No"'], stopIfTrue=True, **warn))
    banding_cf(ws, f"A{first}:W{last}")
    ws.freeze_panes = ws.cell(row=first, column=2)
    ws.auto_filter.ref = f"A{hr}:W{last}"
    print_setup(ws)
    return ws, first, last


def add_redirect_summary(ws, r, first, last):
    S = "'Redirect map'"
    col = lambda c: f"{S}!${c}${first}:${c}${last}"
    r += 2
    ws.cell(row=r, column=2, value="Redirect map").font = font(12, True, C["accent"])
    rows = [
        ("Old URLs listed", f"=COUNTA({col('A')})", None),
        ("High-priority old URLs", f'=COUNTIF({col("D")},"High")', None),
        ("Decision made", f"=COUNTA({col('K')})", None),
        ("Needs decision", f'=COUNTIF({col("T")},"*Needs decision*")', "warn"),
        ("High-priority URLs still undecided", f'=COUNTIFS({col("D")},"High",{col("K")},"")', "warn"),
        ("Accepted suggestion", f'=COUNTIF({col("K")},"Accept suggestion")', None),
        ("Overridden", f'=COUNTIF({col("K")},"Override")', None),
        ("410 or 404", f'=COUNTIF({col("K")},"410 Gone")+COUNTIF({col("K")},"404")', None),
        ("Flag: home page redirect", f'=COUNTIF({col("T")},"*Home page*")', "warn"),
        ("Flag: below your similarity threshold", f'=COUNTIF({col("T")},"*Below your threshold*")', "warn"),
        ("Flag: chain (more than one hop)", f'=COUNTIF({col("T")},"*Chain*")', "warn"),
        ("Flag: destination not 200", f'=COUNTIF({col("T")},"*not 200*")', "warn"),
        ("Flag: redirects to itself", f'=COUNTIF({col("T")},"*to itself*")', "warn"),
        ("QA: approved", f'=COUNTIF({col("U")},"Approved")', None),
        ("QA: live and verified", f'=COUNTIF({col("U")},"Live and verified")', None),
        ("QA: problem", f'=COUNTIF({col("U")},"Problem")', "warn"),
    ]
    start = r + 1
    verified_row = None
    for label, f_, kind in rows:
        r += 1
        if label == "QA: live and verified":
            verified_row = r
        ws.cell(row=r, column=2, value=label).font = font(10, True)
        c = ws.cell(row=r, column=3, value=f_)
        c.font = font(10)
        c.alignment = Alignment(horizontal="right")
        if kind == "warn":
            ws.conditional_formatting.add(f"C{r}", CellIsRule(operator="greaterThan", formula=["0"], fill=fill(C["warn_fill"]), font=Font(name=FONT, color=C["warn_font"], bold=True)))
    r += 1
    ws.cell(row=r, column=2, value="% of old URLs live and verified").font = font(10, True)
    c = ws.cell(row=r, column=3, value=f"=IFERROR(C{verified_row}/C{start},0)")
    c.number_format = "0%"
    c.font = font(12, True, C["accent"])
    bar = ws.cell(row=r, column=4, value=f'=REPT("█",ROUND(C{r}*20,0))&REPT("░",20-ROUND(C{r}*20,0))')
    bar.font = Font(name=FONT, size=11, color=C["accent"])
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=8)
    ws.conditional_formatting.add(f"B{start}:C{r}", FormulaRule(formula=["MOD(ROW(),2)=0"], fill=fill(C["band"])))
    return r


def main():
    checks = read_checks()
    wb = Workbook()
    build_start(wb, START)
    ws, last_check, last = build_checklist(wb, checks, area_label="Phase")
    _, first, rlast = build_redirect_map(wb)
    prog, r = build_progress(wb, "Checklist", AREAS, area_label="Phase", last_row=last)
    add_redirect_summary(prog, r, first, rlast)
    wb.save(HERE / XLSX)
    write_markdown(HERE / "checklist.md", TITLE.replace(" + Redirect Mapping Workbook", ""), MD_INTRO, checks, AREAS, MD_OUTRO)
    print("wrote", XLSX, "and checklist.md with", len(checks), "checks")


MD_INTRO = """
*43 checks across seven phases, from scoping to the first year after launch. By [Lazarina Stoy](https://www.mlforseo.com/experts/lazarina-stoy/) for [MLforSEO](https://www.mlforseo.com/). Built on Google's site move documentation and MLforSEO's fuzzy-matching redirect workflow. Checked October 2026.*

This is the plain-text version of the checklist in `site-migration-checklist.xlsx`. The spreadsheet adds a **Redirect map** tab: paste the output of MLforSEO's fuzzy-matching notebooks, decide per URL, and it builds final destinations, source paths, Apache rules and warning flags, then tracks QA for every URL.

Each check has an **evidence level**:

| Evidence | Meaning |
|---|---|
| Google-documented | Stated in Google Search Central or Google Help documentation |
| Platform-documented | Stated in another platform's own docs (Microsoft Bing, IndexNow) |
| MLforSEO practice | What we teach and use on migrations; sensible and low-risk, but not confirmed by the platforms |

**Priority:** High = skipping it risks lost traffic, indexing or link value. Medium = reduces risk and speeds recovery. Low = tidy-up and documentation.
"""

MD_OUTRO = """
---

## Redirect mapping tools

| Situation | Tool |
|---|---|
| Under about 300 URLs | [Fuzzy Matching for SEO Google Sheets template](https://www.mlforseo.com/resource/fuzzy-matching-for-seo-google-sheets-template/) |
| Old crawl vs new crawl (URL, title, H1, H2, meta description) | [Triple fuzzy matching redirect notebook](https://www.mlforseo.com/resource/automated-redirect-mapping-with-triple-fuzzy-matching-for-site-migrations-notebook/) |
| 404s to live pages | [Mapping 404 URLs to live pages notebook](https://www.mlforseo.com/resource/mapping-404-urls-to-live-pages-with-fuzzy-matching-guide/) |
| Python on GitHub | [`fuzzy-matching-for-seo/`](../fuzzy-matching-for-seo/) |

Fuzzy matching compares characters, not meaning, so review every partial match before it goes live. The walkthrough is in [404 and redirect mapping with fuzzy matching](https://www.mlforseo.com/machine-learning-tutorials/fuzzy-matching/404-and-redirect-mapping-with-fuzzy-matching-in-google-colab-python/).

## Go deeper

- [Google: how to move a site with URL changes](https://developers.google.com/search/docs/crawling-indexing/site-move-with-url-changes)
- [Introduction to Machine Learning for SEO](https://www.mlforseo.com/courses/introduction-to-machine-learning-for-seo/) (fuzzy matching module, MLforSEO Academy)

Created by Lazarina Stoy · MLforSEO. Free to use and adapt under the repository's MIT licence; please keep the credit.
"""

if __name__ == "__main__":
    main()
