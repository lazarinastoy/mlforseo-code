#!/usr/bin/env python3
"""Build topical-map-eav-template.xlsx (MLforSEO, Lazarina Stoy).

    pip install openpyxl
    python build_template.py

You describe one central entity, its attributes and their values (the Entity-Attribute-Variable / EAV model).
The "Topical Map" tab then generates pillar, cluster and supporting rows with titles, slugs, internal-link
suggestions, intent, format, coverage and priority. Plain INDEX/MATCH, COUNTIFS, SUMIFS, SUMPRODUCT and IF
formulas, so it behaves the same in Google Sheets and Excel.
"""
import os
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, PatternFill

from brand_style import (DARK, F_DARK, F_GREEN, F_HI, F_LAV, F_RED, band, credits, esc, flag_rule, font, group_label,
                         header_row, legend, lines, list_validation, note, number_validation, sample_banner, title,
                         widths)

OUT = Path(os.environ.get("OUT", Path(__file__).with_name("topical-map-eav-template.xlsx")))
A_FIRST, A_LAST = 3, 3 + int(os.environ.get("N_ATTR", 50)) - 1        # attribute rows
V_FIRST, V_LAST = 3, 3 + int(os.environ.get("N_VAL", 500)) - 1        # value rows
M_FIRST = 3
M_LAST = M_FIRST + 1 + (A_LAST - A_FIRST + 1) + (V_LAST - V_FIRST + 1) - 1   # map rows
INTENTS = ["Informational", "Commercial", "Transactional", "Navigational"]
EAV_POST = "https://www.mlforseo.com/ml-fundamentals/what-is-the-eav-model-in-seo-a-practical-framework-for-semantic-keyword-research/"
KU_POST = "https://www.mlforseo.com/machine-learning-implementation-guides/keyword-research/what-does-a-good-semantic-keyword-universe-actually-look-like/"
KG_POST = "https://www.mlforseo.com/ml-fundamentals/what-is-a-knowledge-graph-how-to-use-it-for-smarter-keyword-strategy/"
COURSE_URL = "https://www.mlforseo.com/courses/semantic-ml-enabled-keyword-research/"

S = "https://www.stride-example.com"
ENTITY = [
    ("Central entity", "Running shoes", "The one entity the map is about: your main product category, service or topic."),
    ("Entity type", "Product", "Person, Place, Organisation, Product, Event, Creative work or Concept."),
    ("Site root URL", S + "/", "Used only for reference."),
    ("Pillar page URL (if it exists)", S + "/running-shoes/", "Leave blank if the pillar is still a gap."),
    ("Pillar title pattern", "{entity}: the complete guide", "{entity} is replaced by the central entity."),
    ("Cluster title pattern (default)", "{entity} by {attribute}", "Used when an attribute has no cluster title of its own."),
    ("Supporting title pattern (default)", "{entity} for {value}", "Used when an attribute has no pattern of its own."),
    ("Audience / notes", "Recreational and club runners in the UK", ""),
]
# attribute, type, cutoff rule, intent, relevance, include, cluster title, supporting pattern, hub URL, notes
ATTRS = [
    ("Terrain", "Finite", "", "Commercial", 3, "Yes", "", "{value} {entity}", "", ""),
    ("Distance", "Finite", "", "Commercial", 3, "Yes", "", "{entity} for {value}", "", ""),
    ("Foot type", "Finite", "", "Commercial", 3, "Yes", "Running shoes for every foot type", "{entity} for {value}", "", ""),
    ("Cushioning", "Finite", "", "Informational", 2, "Yes", "", "{value} {entity}", "", ""),
    ("Brand", "Finite with a cutoff", "Brands we stock (12)", "Commercial", 2, "Yes", "Running shoe brands", "{value} {entity}", S + "/brands/", ""),
    ("Runner level", "Finite", "", "Informational", 2, "Yes", "", "{entity} for {value}", "", ""),
    ("Price", "Finite with a cutoff", "Three bands: under £80, £80-£140, over £140", "Transactional", 1, "Yes", "", "{entity} {value}", "", ""),
    ("Runner name", "Infinite", "", "Navigational", 1, "No", "", "", "", "Excluded: a page per runner has no reason to exist"),
]
# attribute, value, is entity, user benefit, volume, existing url, notes
VALUES = [
    ("Terrain", "road", "No", "Yes", 6600, S + "/road-running-shoes/", ""),
    ("Terrain", "trail", "No", "Yes", 9900, S + "/trail-running-shoes/", ""),
    ("Terrain", "track", "No", "Yes", 320, "", ""),
    ("Terrain", "treadmill", "No", "Unsure", 590, "", "Check the SERP: mostly treadmill reviews?"),
    ("Distance", "5K", "No", "Yes", 210, "", ""),
    ("Distance", "half marathon", "No", "Yes", 480, "", ""),
    ("Distance", "marathon", "No", "Yes", 2400, S + "/marathon-running-shoes/", ""),
    ("Distance", "ultramarathon", "No", "Yes", 390, "", ""),
    ("Foot type", "wide feet", "No", "Yes", 1900, S + "/wide-fit-running-shoes/", ""),
    ("Foot type", "narrow feet", "No", "Yes", 590, "", ""),
    ("Foot type", "flat feet", "No", "Yes", 1600, "", ""),
    ("Foot type", "high arches", "No", "Yes", 880, "", ""),
    ("Foot type", "overpronation", "Yes", "Yes", 1300, "", "A concept entity in its own right"),
    ("Cushioning", "minimalist", "No", "Yes", 720, "", ""),
    ("Cushioning", "max-cushion", "No", "Yes", 1000, "", ""),
    ("Cushioning", "carbon-plate", "Yes", "Yes", 1000, "", "Technology entity: explain what it is"),
    ("Brand", "Hoka", "Yes", "Yes", 12100, S + "/brands/hoka/", ""),
    ("Brand", "Brooks", "Yes", "Yes", 8100, "", ""),
    ("Brand", "Asics", "Yes", "Yes", 9900, S + "/brands/asics/", ""),
    ("Brand", "Saucony", "Yes", "Yes", 4400, "", ""),
    ("Brand", "On", "Yes", "Yes", 6600, "", ""),
    ("Runner level", "beginners", "No", "Yes", 1300, "", ""),
    ("Runner level", "returning from injury", "No", "Yes", 90, "", "Low volume, real need"),
    ("Runner level", "intermediate runners", "No", "Unsure", 50, "", ""),
    ("Price", "under £80", "No", "Yes", 880, S + "/sale/", ""),
    ("Price", "£80 to £140", "No", "No", 20, "", "Nobody searches this band"),
    ("Price", "over £140", "No", "Unsure", 40, "", ""),
    ("Runner name", "Eliud Kipchoge", "Yes", "No", 2900, "", "Infinite attribute, excluded"),
]
FORMATS = [("Informational", "Guide or explainer"), ("Commercial", "Comparison or best-of list"),
           ("Transactional", "Category or product listing"), ("Navigational", "Landing page")]


def build():
    wb = Workbook()
    start = wb.active
    start.title = "Start here"
    en = wb.create_sheet("Entity")
    at = wb.create_sheet("Attributes")
    va = wb.create_sheet("Values")
    tm = wb.create_sheet("Topical Map")
    co = wb.create_sheet("Combinations")
    da = wb.create_sheet("Dashboard")
    li = wb.create_sheet("Lists")

    AR = lambda c: f"Attributes!${c}${A_FIRST}:${c}${A_LAST}"
    VR = lambda c: f"Values!${c}${V_FIRST}:${c}${V_LAST}"
    MR = lambda c: f"${c}${M_FIRST}:${c}${M_LAST}"
    NA, NV, ENT = "Entity!$B$14", "Entity!$B$15", "Entity!$B$3"

    # ============================================================ Lists
    title(li, "A1", "Lists used by dropdowns and formats")
    header_row(li, 2, ["Intent", "Suggested format for a supporting page"], 1, "input")
    for i, (k, v) in enumerate(FORMATS, start=3):
        li.cell(row=i, column=1, value=k); c = li.cell(row=i, column=2, value=v); c.fill = F_HI
    widths(li, {"A": 18, "B": 40})

    # ============================================================ Entity
    title(en, "A1", "1. Entity: start with one central entity")
    note(en, "A2", "Yellow cells are yours. Patterns use {entity}, {attribute} and {value} as placeholders.")
    for i, (k, v, why) in enumerate(ENTITY, start=3):
        en.cell(row=i, column=1, value=k).font = font(10, True)
        c = en.cell(row=i, column=2, value=v); c.fill = F_HI; c.font = font(10)
        en.cell(row=i, column=3, value=why).font = font(9, italic=True, color="6B6880")
    list_validation(en, "B4", '"Person,Place,Organisation,Product,Event,Creative work,Concept"')
    en["A13"] = "Calculated"; en["A13"].font = font(11, True, "8448EF")
    en["A14"] = "Attributes included"; en["B14"] = f"=MAX({AR('P')})"
    en["A15"] = "Values kept"; en["B15"] = f"=MAX({VR('J')})"
    en["A16"] = "Rows on the topical map"; en["B16"] = "=1+B14+B15"
    for r in (14, 15, 16):
        en[f"B{r}"].font = font(11, True, DARK)
    widths(en, {"A": 34, "B": 46, "C": 80})
    en.sheet_view.showGridLines = False

    # ============================================================ Attributes
    group_label(at, 1, 1, 10, "2. ATTRIBUTES: the ways people describe or filter the entity", "input")
    group_label(at, 1, 11, 16, "CALCULATED", "calc")
    header_row(at, 2, ["Attribute", "Attribute type", "Cutoff rule", "Typical intent", "Business relevance (1-3)", "Include?",
                       "Cluster page title (optional)", "Supporting title pattern (optional)", "Cluster / hub page URL (if it exists)",
                       "Notes"], 1, "input")
    header_row(at, 2, ["Values listed", "Values kept", "Values covered", "Coverage", "Check", "Map order (helper)"], 11, "calc")
    for r in range(A_FIRST, A_LAST + 1):
        e = esc(f"A{r}")
        at[f"K{r}"] = f'=IF(A{r}="","",COUNTIF({VR("A")},{e}))'
        at[f"L{r}"] = f'=IF(A{r}="","",COUNTIFS({VR("A")},{e},{VR("I")},"Yes"))'
        at[f"M{r}"] = f'=IF(A{r}="","",SUMPRODUCT(({VR("A")}=A{r})*({VR("I")}="Yes")*(LEN({VR("F")})>0)))'
        at[f"N{r}"] = f'=IF(A{r}="","",IF(L{r}=0,"",M{r}/L{r}))'
        at[f"O{r}"] = (f'=IF(A{r}="","",IF(AND(B{r}="Infinite",F{r}="Yes"),"Infinite attribute included: add a cutoff or exclude it",'
                       f'IF(AND(B{r}="Finite with a cutoff",C{r}=""),"Add the cutoff rule",IF(AND(F{r}="Yes",L{r}=0),"No values kept yet",""))))')
        at[f"P{r}"] = f'=IF(AND(A{r}<>"",F{r}="Yes"),SUMPRODUCT(($F${A_FIRST}:F{r}="Yes")*(LEN($A${A_FIRST}:A{r})>0)),"")'
        at[f"N{r}"].number_format = "0%"
    for i, row in enumerate(ATTRS):
        for j, v in enumerate(row):
            if v != "":
                at.cell(row=A_FIRST + i, column=1 + j, value=v)
    sample_banner(at, "K1", "SAMPLE: fictional running shoe shop. Clear A-J and add your own.")
    list_validation(at, f"B{A_FIRST}:B{A_LAST}", '"Finite,Finite with a cutoff,Infinite"')
    list_validation(at, f"D{A_FIRST}:D{A_LAST}", '"' + ",".join(INTENTS) + '"')
    list_validation(at, f"E{A_FIRST}:E{A_LAST}", '"1,2,3"')
    list_validation(at, f"F{A_FIRST}:F{A_LAST}", '"Yes,No"')
    band(at, f"A{A_FIRST}:J{A_LAST}", "A", A_FIRST)
    flag_rule(at, f"O{A_FIRST}:O{A_LAST}", f'$O{A_FIRST}<>""', F_HI, bold=True)
    flag_rule(at, f"F{A_FIRST}:F{A_LAST}", f'$F{A_FIRST}="No"', F_RED)
    widths(at, {"A": 18, "B": 20, "C": 30, "D": 15, "E": 12, "F": 9, "G": 30, "H": 24, "I": 36, "J": 34,
                "K": 9, "L": 9, "M": 9, "N": 10, "O": 44, "P": 9})
    at.freeze_panes = "B3"
    at.auto_filter.ref = f"A2:P{A_LAST}"

    # ============================================================ Values
    group_label(va, 1, 1, 7, "3. VALUES: the specific values each attribute takes (sub-entities welcome)", "input")
    group_label(va, 1, 8, 15, "CALCULATED", "calc")
    header_row(va, 2, ["Attribute", "Value / sub-entity", "Is it an entity itself?", "Real user benefit?", "Search volume (optional)",
                       "Existing URL (coverage)", "Notes"], 1, "input")
    header_row(va, 2, ["Attribute included?", "Keep?", "Map order (helper)", "Position in attribute", "Key (helper)",
                       "Kept in attribute", "Sibling value", "Sibling map order (helper)"], 8, "calc")
    for r in range(V_FIRST, V_LAST + 1):
        e = esc(f"A{r}")
        nxt = f'A{r}&"|"&IF(K{r}=M{r},1,K{r}+1)'
        va[f"H{r}"] = f'=IF(A{r}="","",IFERROR(INDEX({AR("F")},MATCH(A{r},{AR("A")},0))&"","Unknown attribute"))'
        va[f"I{r}"] = f'=IF(OR(A{r}="",B{r}=""),"",IF(AND(H{r}="Yes",D{r}<>"No"),"Yes","No"))'
        va[f"J{r}"] = f'=IF(I{r}="Yes",COUNTIF($I${V_FIRST}:I{r},"Yes"),"")'
        va[f"K{r}"] = f'=IF(I{r}<>"Yes","",COUNTIFS($A${V_FIRST}:A{r},{e},$I${V_FIRST}:I{r},"Yes"))'
        va[f"L{r}"] = f'=IF(K{r}="","",A{r}&"|"&K{r})'
        va[f"M{r}"] = f'=IF(I{r}<>"Yes","",COUNTIFS({VR("A")},{e},{VR("I")},"Yes"))'
        va[f"N{r}"] = f'=IF(I{r}<>"Yes","",IF(M{r}<2,"",INDEX({VR("B")},MATCH({nxt},{VR("L")},0))))'
        va[f"O{r}"] = f'=IF(N{r}="","",INDEX({VR("J")},MATCH({nxt},{VR("L")},0)))'
    for i, row in enumerate(VALUES):
        for j, v in enumerate(row):
            if v != "":
                va.cell(row=V_FIRST + i, column=1 + j, value=v)
    sample_banner(va, "H1", "SAMPLE values. Clear A-G and add your own (one value per row).")
    list_validation(va, f"A{V_FIRST}:A{V_LAST}", f"=Attributes!$A${A_FIRST}:$A${A_LAST}")
    list_validation(va, f"C{V_FIRST}:C{V_LAST}", '"Yes,No"')
    list_validation(va, f"D{V_FIRST}:D{V_LAST}", '"Yes,Unsure,No"',
                    prompt="Would a real person benefit from a page about this value? 'No' drops it from the map.")
    number_validation(va, f"E{V_FIRST}:E{V_LAST}", 0)
    band(va, f"A{V_FIRST}:G{V_LAST}", "A", V_FIRST)
    flag_rule(va, f"I{V_FIRST}:I{V_LAST}", f'$I{V_FIRST}="No"', F_RED)
    flag_rule(va, f"H{V_FIRST}:H{V_LAST}", f'$H{V_FIRST}="Unknown attribute"', F_HI, bold=True)
    widths(va, {"A": 18, "B": 26, "C": 13, "D": 13, "E": 12, "F": 40, "G": 36, "H": 12, "I": 8, "J": 9, "K": 9,
                "L": 20, "M": 9, "N": 20, "O": 9})
    va.freeze_panes = "C3"
    va.auto_filter.ref = f"A2:O{V_LAST}"

    # ============================================================ Topical Map (generated)
    group_label(tm, 1, 1, 15, "4. TOPICAL MAP: generated from Entity, Attributes and Values. Don't type here; copy it as values when you're happy.", "calc")
    header_row(tm, 2, ["#", "Level", "Attribute", "Value / sub-entity", "Suggested title", "Suggested slug", "Links up to",
                       "Links across to", "Links down to", "Primary intent", "Suggested format", "Existing URL", "Coverage",
                       "Search volume", "Priority"], 1, "calc")
    header_row(tm, 2, ["Value row (helper)", "Attribute row (helper)", "Raw title (helper)"], 16, "calc")
    for r in range(M_FIRST, M_LAST + 1):
        tm[f"A{r}"] = "=ROW()-2"
        tm[f"B{r}"] = f'=IF(A{r}=1,"Pillar",IF(A{r}<=1+{NA},"Cluster",IF(A{r}<=1+{NA}+{NV},"Supporting","")))'
        tm[f"P{r}"] = f'=IF(B{r}="Supporting",MATCH(A{r}-1-{NA},{VR("J")},0),"")'
        tm[f"C{r}"] = (f'=IF(B{r}="Cluster",INDEX({AR("A")},MATCH(A{r}-1,{AR("P")},0)),'
                       f'IF(B{r}="Supporting",INDEX({VR("A")},P{r}),""))')
        tm[f"Q{r}"] = f'=IF(C{r}="","",MATCH(C{r},{AR("A")},0))'
        tm[f"D{r}"] = f'=IF(B{r}="Supporting",INDEX({VR("B")},P{r}),"")'
        ent = f"LOWER({ENT})"
        cluster_t = (f'IF(INDEX({AR("G")},Q{r})<>"",INDEX({AR("G")},Q{r}),'
                     f'SUBSTITUTE(SUBSTITUTE(Entity!$B$8,"{{entity}}",{ent}),"{{attribute}}",LOWER(C{r})))')
        pattern = f'IF(INDEX({AR("H")},Q{r})<>"",INDEX({AR("H")},Q{r}),Entity!$B$9)'
        support_t = f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({pattern},"{{entity}}",{ent}),"{{value}}",D{r}),"{{attribute}}",LOWER(C{r}))'
        tm[f"R{r}"] = (f'=IF(B{r}="Pillar",SUBSTITUTE(Entity!$B$7,"{{entity}}",{ent}),IF(B{r}="Cluster",{cluster_t},'
                       f'IF(B{r}="Supporting",{support_t},"")))')
        tm[f"E{r}"] = f'=IF(R{r}="","",UPPER(LEFT(R{r},1))&MID(R{r},2,500))'
        tm[f"F{r}"] = (f'=IF(E{r}="","",LOWER(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(TRIM(E{r}),":",""),",",""),'
                       f'"\'",""),"&","and"),"£",""),"  "," ")," ","-")))')
        tm[f"G{r}"] = f'=IF(B{r}="Cluster",$E${M_FIRST},IF(B{r}="Supporting",INDEX({MR("E")},MATCH(C{r},{MR("C")},0)),""))'
        tm[f"H{r}"] = (f'=IF(B{r}="Supporting",IFERROR(INDEX({MR("E")},1+{NA}+INDEX({VR("O")},P{r})),""),'
                       f'IF(B{r}="Cluster",IF({NA}<2,"",INDEX({MR("E")},IF(A{r}-1={NA},2,A{r}+1))),""))')
        tm[f"I{r}"] = (f'=IF(B{r}="Pillar",{NA}&" cluster pages",IF(B{r}="Cluster",'
                       f'(COUNTIF({MR("C")},{esc(f"C{r}")})-1)&" supporting pages",""))')
        tm[f"J{r}"] = f'=IF(B{r}="Pillar","Informational",IF(OR(B{r}="Cluster",B{r}="Supporting"),INDEX({AR("D")},Q{r})&"",""))'
        tm[f"K{r}"] = (f'=IF(B{r}="","",IF(B{r}="Pillar","Pillar guide (entity hub)",IF(B{r}="Cluster","Hub page for this attribute",'
                       f'IF(INDEX({VR("C")},P{r})="Yes","Entity page (profile of the sub-entity)",'
                       f'IFERROR(INDEX(Lists!$B$3:$B$6,MATCH(J{r},Lists!$A$3:$A$6,0)),"Article")))))')
        tm[f"L{r}"] = (f'=IF(B{r}="Pillar",Entity!$B$6&"",IF(B{r}="Cluster",INDEX({AR("I")},Q{r})&"",'
                       f'IF(B{r}="Supporting",INDEX({VR("F")},P{r})&"","")))')
        tm[f"M{r}"] = f'=IF(B{r}="","",IF(L{r}<>"","Covered","Gap"))'
        tm[f"N{r}"] = (f'=IF(B{r}="Pillar",SUMIFS({VR("E")},{VR("I")},"Yes"),IF(B{r}="Cluster",SUMIFS({VR("E")},{VR("A")},{esc(f"C{r}")},{VR("I")},"Yes"),'
                       f'IF(B{r}="Supporting",N(INDEX({VR("E")},P{r})),"")))')
        tm[f"O{r}"] = (f'=IF(B{r}="","",IF(B{r}="Pillar",100,IF(B{r}="Cluster",N(INDEX({AR("E")},Q{r}))*25+IF(M{r}="Gap",10,0),'
                       f'N(INDEX({AR("E")},Q{r}))*20+MIN(20,ROUND(LOG10(N(N{r})+1)*5,0))+IF(M{r}="Gap",10,0)'
                       f'+IF(INDEX({VR("D")},P{r})="Yes",5,0))))')
    tm.conditional_formatting  # noqa
    flag_rule(tm, f"A{M_FIRST}:O{M_LAST}", f'$B{M_FIRST}="Pillar"', PatternFill("solid", fgColor="D9CCF9"), bold=True)
    flag_rule(tm, f"A{M_FIRST}:O{M_LAST}", f'$B{M_FIRST}="Cluster"', F_LAV, bold=True)
    flag_rule(tm, f"M{M_FIRST}:M{M_LAST}", f'$M{M_FIRST}="Gap"', F_HI, bold=True)
    flag_rule(tm, f"M{M_FIRST}:M{M_LAST}", f'$M{M_FIRST}="Covered"', F_GREEN)
    widths(tm, {"A": 5, "B": 11, "C": 14, "D": 22, "E": 42, "F": 36, "G": 34, "H": 34, "I": 20, "J": 14, "K": 30,
                "L": 40, "M": 10, "N": 11, "O": 9, "P": 9, "Q": 9, "R": 9})
    tm.freeze_panes = "C3"
    tm.auto_filter.ref = f"A2:O{M_LAST}"

    # ============================================================ Combinations
    title(co, "A1", "5. Combinations: two attributes crossed (long-tail ideas)")
    note(co, "A2", "Pick two finite attributes. Copy the combinations worth a page into your brief or back into Values.")
    for i, (k, v) in enumerate([("Attribute A (rows)", "Terrain"), ("Attribute B (columns)", "Foot type"),
                                ("Title pattern", "{a} {entity} for {b}")], start=3):
        co.cell(row=i, column=1, value=k).font = font(10, True)
        c = co.cell(row=i, column=2, value=v); c.fill = F_HI
    list_validation(co, "B3:B4", f"=Attributes!$A${A_FIRST}:$A${A_LAST}")
    co["A6"] = (f'=IF(OR(IFERROR(INDEX({AR("B")},MATCH(B3,{AR("A")},0)),"")="Infinite",'
                f'IFERROR(INDEX({AR("B")},MATCH(B4,{AR("A")},0)),"")="Infinite"),'
                f'"Warning: an infinite attribute never runs out of combinations. Pick finite ones.","")')
    co["A6"].font = font(10, True, "B42318")
    co["A8"] = "A \\ B"
    co["A8"].fill = F_DARK; co["A8"].font = font(10, True, "FFFFFF")
    cols = 10
    rows = 15
    from openpyxl.utils import get_column_letter as L
    for j in range(1, cols + 1):
        c = co.cell(row=8, column=1 + j, value=f'=IFERROR(INDEX({VR("B")},MATCH($B$4&"|"&{j},{VR("L")},0)),"")')
        c.fill = F_DARK; c.font = font(10, True, "FFFFFF"); c.alignment = Alignment(wrap_text=True)
        co.column_dimensions[L(1 + j)].width = 26
    for i in range(1, rows + 1):
        r = 8 + i
        c = co.cell(row=r, column=1, value=f'=IFERROR(INDEX({VR("B")},MATCH($B$3&"|"&{i},{VR("L")},0)),"")')
        c.fill = F_LAV; c.font = font(10, True)
        for j in range(1, cols + 1):
            cl = L(1 + j)
            t = f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE($B$5,"{{a}}",$A{r}),"{{b}}",{cl}$8),"{{entity}}",LOWER({ENT}))'
            c = co.cell(row=r, column=1 + j, value=f'=IF(OR($A{r}="",{cl}$8=""),"",UPPER(LEFT({t},1))&MID({t},2,300))')
            c.alignment = Alignment(wrap_text=True, vertical="top"); c.font = font(9)
    co.column_dimensions["A"].width = 24
    co.freeze_panes = "B9"

    # ============================================================ Dashboard
    title(da, "A1", "Topical map dashboard")
    note(da, "A2", "Coverage = kept values that already have a URL. Gaps are your content plan.")
    header_row(da, 3, ["Metric", "Value"], 1, "calc")
    kpis = [
        ("Central entity", f"={ENT}"),
        ("Attributes listed", f"=SUMPRODUCT((LEN({AR('A')})>0)*1)"),
        ("Attributes included", f"={NA}"),
        ("Values listed", f"=SUMPRODUCT((LEN({VR('B')})>0)*1)"),
        ("Values kept on the map", f"={NV}"),
        ("Rows on the topical map", "=Entity!$B$16"),
        ("Covered (has a URL)", f'=COUNTIF(\'Topical Map\'!{MR("M")},"Covered")'),
        ("Gaps (no URL yet)", f'=COUNTIF(\'Topical Map\'!{MR("M")},"Gap")'),
        ("Coverage", '=IF(B10+B11=0,"",B10/(B10+B11))'),
        ("Attribute checks to resolve", f'=SUMPRODUCT((LEN({AR("O")})>0)*1)'),
    ]
    for i, (k, v) in enumerate(kpis, start=4):
        da.cell(row=i, column=1, value=k).font = font(10)
        c = da.cell(row=i, column=2, value=v); c.font = font(12, True, DARK); c.alignment = Alignment(horizontal="center")
    da["B12"].number_format = "0%"
    flag_rule(da, "B13", "AND(ISNUMBER(B13),B13>0)", F_HI)
    header_row(da, 3, ["Attribute", "Included?", "Values kept", "Covered", "Gaps", "Coverage", "Check"], 4, "calc")
    for i in range(A_LAST - A_FIRST + 1):
        r = 4 + i
        src = A_FIRST + i
        da.cell(row=r, column=4, value=f'=IF(Attributes!A{src}="","",Attributes!A{src})')
        da.cell(row=r, column=5, value=f'=IF(D{r}="","",Attributes!F{src}&"")')
        da.cell(row=r, column=6, value=f'=IF(D{r}="","",Attributes!L{src})')
        da.cell(row=r, column=7, value=f'=IF(D{r}="","",Attributes!M{src})')
        da.cell(row=r, column=8, value=f'=IF(D{r}="","",F{r}-G{r})')
        c = da.cell(row=r, column=9, value=f'=IF(D{r}="","",Attributes!N{src})'); c.number_format = "0%"
        da.cell(row=r, column=10, value=f'=IF(D{r}="","",Attributes!O{src}&"")')
    flag_rule(da, f"J4:J{4 + A_LAST - A_FIRST}", 'LEN($J4)>0', F_HI)
    widths(da, {"A": 30, "B": 22, "C": 3, "D": 20, "E": 10, "F": 11, "G": 10, "H": 8, "I": 10, "J": 50})
    da.sheet_view.showGridLines = False

    # ============================================================ Start here
    s = start
    title(s, "A1", "Topical Map Template + Generator (EAV-based)")
    s["A2"] = "Entity → attributes → values → pillar, cluster and supporting pages · by Lazarina Stoy · MLforSEO"
    s["A2"].font = font(11, False, "8448EF", True)
    r = lines(s, 4, [
        "What it does",
        "You describe one central entity, the attributes people use to describe or filter it, and the values each attribute takes.",
        "The Topical Map tab turns that into a pillar page, one cluster (hub) page per attribute and one supporting page per value,",
        "with a suggested title, slug, intent, format, internal links (up, across and down), coverage against your existing URLs",
        "and a priority score. Combinations crosses two attributes for long-tail ideas.",
        "",
        "EAV (entity, attribute, variable or value) is a data-modelling pattern borrowed from database design. I use it here as a",
        "planning lens: a structured way to list what your audience wants to know about an entity and to spot gaps. It is not a",
        "claim about how any search engine ranks pages.",
    ], bold_first=True)
    r = lines(s, r + 1, [
        "How to use it",
        "1. Entity: name the central entity and adjust the three title patterns ({entity}, {attribute}, {value}).",
        "2. Attributes: list every attribute your audience cares about. Mark each Finite, Finite with a cutoff, or Infinite.",
        "   Infinite attributes (names, URLs, slogans) need a cutoff or should be excluded. Set relevance 1-3 and Include? Yes/No.",
        "3. Values: one row per value. Mark sub-entities, and answer 'Real user benefit?' honestly: 'No' drops the value.",
        "   Add the existing URL where you already cover it. Volumes are optional (autocomplete or Knowledge Graph notebooks help).",
        "4. Topical Map: read it top to bottom. Gaps are yellow. Sort or filter by Priority to plan the next quarter.",
        "5. When you're happy, copy the Topical Map and paste it as values into a new tab, so your edits stay put.",
        "6. Combinations: pick two finite attributes and copy the pairs worth a page into your content brief.",
    ], bold_first=True)
    r = lines(s, r + 1, [
        "How the map is built",
        "Pillar: the central entity. Clusters: included attributes, in the order you listed them. Supporting: kept values.",
        "Links up to: supporting pages link to their cluster page; cluster pages link to the pillar. Link every page to the pillar too.",
        "Links across to: the next value in the same attribute (and the next cluster), so siblings link to each other in a loop.",
        "Priority: pillar 100; cluster = relevance x 25 (+10 if it's a gap); supporting = relevance x 20 + volume points (max 20)",
        "   + 10 if it's a gap + 5 if you marked real user benefit 'Yes'. Edit the weights in the formula if you prefer others.",
        "Capacity: 50 attributes and 500 values. Titles and slugs are suggestions, so edit them once you copy the map as values.",
    ], bold_first=True)
    r = legend(s, r + 1)
    r = credits(s, r, "Read: What is the EAV model in SEO? (MLforSEO)", EAV_POST,
                "Course: Semantic ML-enabled Keyword Research (MLforSEO Academy)", COURSE_URL,
                [("Read: What does a good semantic keyword universe look like? (MLforSEO)", KU_POST),
                 ("Read: What is a knowledge graph? How to use it for keyword strategy (MLforSEO)", KG_POST)])
    widths(s, {"A": 30, "B": 110})
    s.sheet_view.showGridLines = False
    s.sheet_properties.tabColor = DARK
    tm.sheet_properties.tabColor = "8448EF"
    da.sheet_properties.tabColor = "FFDE59"
    wb.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    build()
