"""MLforSEO brand styling helpers for openpyxl workbooks (shared by the template builders).

Palette: header #251046 (white text), accent #8448ef, soft lavender #f4f0fd, highlight #ffde59.
Inputs get dark headers, calculated columns get accent (purple) headers.
"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

DARK = "251046"
ACCENT = "8448EF"
LAVENDER = "F4F0FD"
HIGHLIGHT = "FFDE59"
SOFT_RED = "FBE3E4"
SOFT_GREEN = "E3F5EA"
GREY = "6B6880"
FONT = "Inter"

F_DARK = PatternFill("solid", fgColor=DARK)
F_ACCENT = PatternFill("solid", fgColor=ACCENT)
F_LAV = PatternFill("solid", fgColor=LAVENDER)
F_HI = PatternFill("solid", fgColor=HIGHLIGHT)
F_RED = PatternFill("solid", fgColor=SOFT_RED)
F_GREEN = PatternFill("solid", fgColor=SOFT_GREEN)

THIN = Side(style="thin", color="D9D2EE")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def font(size=10, bold=False, color="1A0B32", italic=False, underline=None):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic, underline=underline)


def title(ws, cell, text, size=18):
    ws[cell] = text
    ws[cell].font = font(size, True, DARK)


def note(ws, cell, text, color=GREY, italic=True, size=9):
    ws[cell] = text
    ws[cell].font = font(size, False, color, italic)
    ws[cell].alignment = Alignment(wrap_text=False, vertical="top")


def header_row(ws, row, headers, start_col=1, kind="input", height=34):
    """Write a header row. kind: 'input' (dark) or 'calc' (accent purple)."""
    fill = F_DARK if kind == "input" else F_ACCENT
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=start_col + i, value=h)
        c.fill = fill
        c.font = font(10, True, "FFFFFF")
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BOX
    ws.row_dimensions[row].height = height


def group_label(ws, row, col_from, col_to, text, kind="input"):
    fill = F_LAV if kind == "input" else PatternFill("solid", fgColor="E7DDFD")
    for c in range(col_from, col_to + 1):
        ws.cell(row=row, column=c).fill = fill
    c = ws.cell(row=row, column=col_from, value=text)
    c.font = font(9, True, DARK if kind == "input" else ACCENT)


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


def set_widths_by_index(ws, start_col, ws_widths):
    for i, w in enumerate(ws_widths):
        ws.column_dimensions[get_column_letter(start_col + i)].width = w


def list_validation(ws, sqref, formula1, prompt=None, strict=True):
    dv = DataValidation(type="list", formula1=formula1, allow_blank=True)
    dv.showErrorMessage = True
    dv.errorStyle = "stop" if strict else "warning"
    if prompt:
        dv.promptTitle = "Pick a value"
        dv.prompt = prompt
        dv.showInputMessage = True
    dv.add(sqref)
    ws.add_data_validation(dv)
    return dv


def number_validation(ws, sqref, lo=0, hi=None, decimal=True):
    if hi is None:
        dv = DataValidation(type="decimal" if decimal else "whole", operator="greaterThanOrEqual", formula1=str(lo), allow_blank=True)
    else:
        dv = DataValidation(type="decimal" if decimal else "whole", operator="between", formula1=str(lo), formula2=str(hi), allow_blank=True)
    dv.errorStyle = "warning"
    dv.showErrorMessage = True
    dv.add(sqref)
    ws.add_data_validation(dv)


def band(ws, rng, first_col_letter, first_row):
    """Lavender banding on every other filled row (conditional, so it survives pasting)."""
    ws.conditional_formatting.add(
        rng, FormulaRule(formula=[f'AND(${first_col_letter}{first_row}<>"",MOD(ROW(),2)=0)'], fill=F_LAV))


def flag_rule(ws, rng, formula, fill=F_HI, bold=False):
    ws.conditional_formatting.add(rng, FormulaRule(formula=[formula], fill=fill, font=Font(bold=bold) if bold else None))


def credits(ws, row, post_text, post_url, course_text, course_url, extra=None):
    ws.cell(row=row, column=1, value="Credits").font = font(12, True, DARK)
    ws.cell(row=row + 1, column=1, value="Template by Lazarina Stoy · MLforSEO (mlforseo.com). Free to use and adapt; please keep this credit.").font = font(10)
    c = ws.cell(row=row + 2, column=1, value=post_text)
    c.hyperlink = post_url
    c.font = font(10, False, ACCENT, underline="single")
    c = ws.cell(row=row + 3, column=1, value=course_text)
    c.hyperlink = course_url
    c.font = font(10, False, ACCENT, underline="single")
    r = row + 4
    for text, url in (extra or []):
        c = ws.cell(row=r, column=1, value=text)
        c.hyperlink = url
        c.font = font(10, False, ACCENT, underline="single")
        r += 1
    return r


def legend(ws, row):
    ws.cell(row=row, column=1, value="Legend").font = font(12, True, DARK)
    items = [
        (F_DARK, "FFFFFF", "Dark header", "Input column: paste or type here."),
        (F_ACCENT, "FFFFFF", "Purple header", "Calculated column: formulas, please don't type over them."),
        (F_LAV, DARK, "Lavender band", "Readability banding and paste zones."),
        (F_HI, DARK, "Yellow cell", "A flag that needs a human look, or an editable setting."),
        (F_RED, DARK, "Soft red", "A stronger warning (for example cannibalisation or a prune decision)."),
    ]
    for i, (fill, fc, name, text) in enumerate(items, start=1):
        a = ws.cell(row=row + i, column=1, value=name)
        a.fill = fill
        a.font = font(10, True, fc)
        ws.cell(row=row + i, column=2, value=text).font = font(10)
    return row + len(items) + 1


def lines(ws, row, texts, col=1, size=10, bold_first=False):
    for i, t in enumerate(texts):
        c = ws.cell(row=row + i, column=col, value=t)
        c.font = font(size, bold_first and i == 0)
        c.alignment = Alignment(wrap_text=False, vertical="top")
    return row + len(texts)


def sample_banner(ws, cell, text="SAMPLE DATA (synthetic, for a fictional site). Delete these rows and paste your own."):
    ws[cell] = text
    ws[cell].fill = F_HI
    ws[cell].font = font(9, True, DARK)


def esc(ref):
    """Escape ? and * so COUNTIFS/SUMIFS treat a URL or keyword literally."""
    return f'SUBSTITUTE(SUBSTITUTE({ref},"?","~?"),"*","~*")'
