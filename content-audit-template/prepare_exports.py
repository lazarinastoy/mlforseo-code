#!/usr/bin/env python3
"""Normalise Search Console, GA4 and Screaming Frog exports into the paste format of content-audit-template.xlsx.

Optional helper: you can paste the raw exports into the workbook by hand. This script saves the fiddly bits:
it strips GA4's "#" comment lines, reads the Search Console zip, renames common header variants to the names the
workbook expects, reports how well the three sources join on URL, and (with --fill) writes everything straight into
a copy of the workbook.

    pip install pandas openpyxl
    python prepare_exports.py --gsc Pages.csv --ga4 ga4.csv --crawl internal_html.csv --out prepared/
    python prepare_exports.py --gsc "Performance.zip" --gsc-previous prev.zip --ga4 ga4.csv --crawl internal_html.csv \
        --out prepared/ --fill content-audit-template.xlsx
    python prepare_exports.py --demo          # runs on sample_data/ and fills a copy of the template

By Lazarina Stoy, MLforSEO (MIT licence).
"""
import argparse
import io
import re
import sys
import zipfile
from pathlib import Path

import pandas as pd

HERE = Path(__file__).resolve().parent
MAX_ROWS = 2000  # rows each input tab holds (copy the formula rows down in the sheet if you need more)

GSC_COLUMNS = ["Top pages", "Clicks", "Impressions", "CTR", "Position"]
GSC_ALIASES = {"top pages": "Top pages", "page": "Top pages", "pages": "Top pages", "url": "Top pages", "landing page": "Top pages",
               "clicks": "Clicks", "impressions": "Impressions", "ctr": "CTR", "url ctr": "CTR", "position": "Position",
               "average position": "Position", "avg. position": "Position"}
GA4_COLUMNS = ["Landing page + query string", "Sessions", "Active users", "New users", "Average engagement time per session",
               "Key events", "Total revenue"]
GA4_ALIASES = {"landing page + query string": "Landing page + query string", "landing page": "Landing page + query string",
               "page path and screen class": "Landing page + query string", "page path + query string": "Landing page + query string",
               "sessions": "Sessions", "active users": "Active users", "users": "Active users", "new users": "New users",
               "average engagement time per session": "Average engagement time per session",
               "average engagement time": "Average engagement time per session",
               "user engagement duration per session": "Average engagement time per session",
               "key events": "Key events", "conversions": "Key events", "total revenue": "Total revenue"}
CRAWL_REQUIRED = ["Address", "Status Code", "Indexability", "Indexability Status", "Word Count", "Unique Inlinks"]
CRAWL_OPTIONAL = ["Closest Near Duplicate Match", "No. Near Duplicates", "Closest Semantically Similar Address", "Semantic Similarity Score"]


# ----------------------------------------------------------------------------- URL key (same logic as the sheet)
def url_key(value) -> str:
    """Path without host, query string, fragment or trailing slash, lower-cased. Mirrors the helper columns in the sheet."""
    s = str(value or "").strip()
    if not s:
        return ""
    if s.lower().startswith("http"):
        m = re.match(r"https?://[^/]+(/.*)?$", s, flags=re.I)
        s = (m.group(1) if m and m.group(1) else "/")
    s = re.split(r"[?#]", s, maxsplit=1)[0]
    if len(s) > 1 and s.endswith("/"):
        s = s[:-1]
    return s.lower()


# ----------------------------------------------------------------------------- readers
def _rename(df, aliases):
    return df.rename(columns={c: aliases.get(str(c).strip().lower(), str(c).strip()) for c in df.columns})


def _to_number(series):
    """'1,234' -> 1234, '4.1%' -> 0.041, '1m 05s' -> 65. Leaves blanks blank."""
    def conv(v):
        if pd.isna(v) or str(v).strip() == "":
            return None
        t = str(v).strip().replace(",", "")
        m = re.fullmatch(r"(?:(\d+)m)?\s*(?:(\d+)s)?", t)
        if m and (m.group(1) or m.group(2)):
            return int(m.group(1) or 0) * 60 + int(m.group(2) or 0)
        if t.endswith("%"):
            return float(t[:-1]) / 100
        try:
            return float(t) if "." in t else int(t)
        except ValueError:
            return v
    return series.map(conv)


def read_gsc(path) -> pd.DataFrame:
    """Search Console > Performance > Pages export: the .zip (uses Pages.csv inside) or the CSV itself."""
    path = Path(path)
    if path.suffix.lower() == ".zip":
        with zipfile.ZipFile(path) as z:
            name = next((n for n in z.namelist() if n.lower().endswith("pages.csv")), None)
            if not name:
                sys.exit(f"{path}: no Pages.csv inside the zip. Export from the Pages tab of the Performance report.")
            df = pd.read_csv(io.BytesIO(z.read(name)), dtype=str)
    else:
        df = pd.read_csv(path, dtype=str, encoding="utf-8-sig")
    df = _rename(df, GSC_ALIASES)
    missing = [c for c in ["Top pages", "Clicks", "Impressions", "Position"] if c not in df.columns]
    if missing:
        sys.exit(f"{path}: missing Search Console columns {missing}. Found {list(df.columns)}")
    for c in ("Clicks", "Impressions", "Position"):
        df[c] = _to_number(df[c])
    if "CTR" not in df.columns:
        df["CTR"] = None
    return df[GSC_COLUMNS]


def read_ga4(path) -> pd.DataFrame:
    """GA4 > Reports > Engagement > Landing page > Share this report > Download file > CSV. Skips the '#' header lines."""
    text = Path(path).read_text(encoding="utf-8-sig")
    lines = [ln for ln in text.splitlines() if not ln.startswith("#")]
    while lines and not lines[0].strip():
        lines.pop(0)
    # GA4 sometimes appends a second table after a blank line: keep the first one only.
    if "" in [ln.strip() for ln in lines]:
        lines = lines[: [ln.strip() for ln in lines].index("")]
    df = pd.read_csv(io.StringIO("\n".join(lines)), dtype=str)
    df = _rename(df, GA4_ALIASES)
    if "Average engagement time per session" not in df.columns and any("per active user" in c.lower() for c in df.columns):
        print("  note: GA4 export has engagement time per active user, not per session. Add 'Average engagement time per session' "
              "to the report (Customise report > Metrics) for a like-for-like number.")
    missing = [c for c in ["Landing page + query string", "Sessions"] if c not in df.columns]
    if missing:
        sys.exit(f"{path}: missing GA4 columns {missing}. Found {list(df.columns)}")
    df = df[df["Landing page + query string"].fillna("").str.strip().str.lower() != "grand total"]
    for c in GA4_COLUMNS[1:]:
        if c in df.columns:
            df[c] = _to_number(df[c])
        else:
            df[c] = None
    return df[GA4_COLUMNS]


def read_crawl(path) -> pd.DataFrame:
    """Screaming Frog Internal > HTML (or Internal > All) export. Keeps every column: the sheet finds them by name."""
    raw = Path(path).read_text(encoding="utf-8-sig")
    first = raw.splitlines()[0] if raw else ""
    skip = 1 if first.count(",") == 0 else 0      # very old exports start with a title line such as "Internal - HTML"
    df = pd.read_csv(io.StringIO(raw), dtype=str, skiprows=skip, keep_default_na=False)
    missing = [c for c in CRAWL_REQUIRED if c not in df.columns]
    if missing:
        sys.exit(f"{path}: missing Screaming Frog columns {missing}.")
    absent = [c for c in CRAWL_OPTIONAL if c not in df.columns]
    if absent:
        print(f"  note: crawl has no {absent}. Merge suggestions need near-duplicate and/or semantic similarity "
              "turned on in Screaming Frog (Config > Content > Duplicates; Config > Content > Embeddings) and Crawl Analysis run.")
    if "Content Type" in df.columns:
        df = df[df["Content Type"].str.contains("html", case=False, na=False) | (df["Content Type"] == "")]
    for c in ["Status Code", "Word Count", "Unique Inlinks", "No. Near Duplicates", "Semantic Similarity Score"]:
        if c in df.columns:
            df[c] = _to_number(df[c])
    return df


# ----------------------------------------------------------------------------- join report
def join_report(crawl, gsc, ga4):
    ck = set(crawl["Address"].map(url_key))
    g = gsc.assign(key=gsc["Top pages"].map(url_key))
    a = ga4.assign(key=ga4["Landing page + query string"].map(url_key))
    g_clicks = pd.to_numeric(g["Clicks"], errors="coerce").fillna(0)
    a_sess = pd.to_numeric(a["Sessions"], errors="coerce").fillna(0)
    out = {
        "crawl URLs": len(ck),
        "GSC pages": g["key"].nunique(),
        "GSC clicks matched to the crawl": f"{g_clicks[g['key'].isin(ck)].sum():,.0f} of {g_clicks.sum():,.0f}",
        "GA4 sessions matched to the crawl": f"{a_sess[a['key'].isin(ck)].sum():,.0f} of {a_sess.sum():,.0f}",
    }
    missing = g[(~g["key"].isin(ck)) & (g_clicks > 0)]
    hosts = set(re.findall(r"https?://([^/]+)", " ".join(crawl["Address"].astype(str).head(200))))
    if len(hosts) > 1:
        print(f"  warning: the crawl has more than one host {sorted(hosts)}. The sheet joins on path only, so audit one host at a time.")
    return out, missing["Top pages"].tolist()


# ----------------------------------------------------------------------------- fill the workbook
def fill_workbook(template, out_path, frames):
    """Write each frame into its input tab (header row 1, data from row 2) and clear the sample rows."""
    from openpyxl import load_workbook
    wb = load_workbook(template)
    widths = {"GSC current": 5, "GSC previous": 5, "GA4": 11, "Crawl": 130}
    for tab, df in frames.items():
        ws = wb[tab]
        ncols = widths[tab]
        if len(df) > MAX_ROWS:
            print(f"  warning: {tab} has {len(df)} rows; the template holds {MAX_ROWS}. Only the first {MAX_ROWS} are written. "
                  "Copy the last row of formulas down in the sheet and paste the rest by hand.")
        for row in ws.iter_rows(min_row=1, max_row=MAX_ROWS + 1, max_col=ncols):
            for c in row:
                c.value = None
        cols = list(df.columns)[:ncols]
        for j, h in enumerate(cols, start=1):
            ws.cell(row=1, column=j, value=h)
        for i, rec in enumerate(df[cols].head(MAX_ROWS).itertuples(index=False), start=2):
            for j, v in enumerate(rec, start=1):
                if v is None or (isinstance(v, float) and pd.isna(v)) or v == "":
                    continue
                ws.cell(row=i, column=j, value=v)
    if "Audit" in wb.sheetnames:   # clear sample manual inputs (referring domains, override, notes)
        ws = wb["Audit"]
        for r in range(3, MAX_ROWS + 3):
            for col in ("AC", "AD", "AF"):
                ws[f"{col}{r}"].value = None
    wb.save(out_path)
    return out_path


def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--gsc", help="Search Console Pages export (.zip or .csv), current period")
    p.add_argument("--gsc-previous", help="Same export for the previous period (optional, for decay)")
    p.add_argument("--ga4", help="GA4 landing page CSV")
    p.add_argument("--crawl", help="Screaming Frog internal_html.csv")
    p.add_argument("--out", default="prepared", help="Folder for the cleaned CSVs")
    p.add_argument("--fill", help="Path to content-audit-template.xlsx: writes <out>/content-audit-filled.xlsx")
    p.add_argument("--demo", action="store_true", help="Run on sample_data/")
    a = p.parse_args(argv)
    if a.demo:
        sd = HERE / "sample_data"
        a.gsc, a.gsc_previous, a.ga4, a.crawl = sd / "gsc_pages_current.csv", sd / "gsc_pages_previous.csv", sd / "ga4_landing_pages.csv", sd / "internal_html.csv"
        a.out = a.out if a.out != "prepared" else str(HERE / "prepared_demo")
        a.fill = a.fill or str(HERE / "content-audit-template.xlsx")
    if not (a.gsc and a.ga4 and a.crawl):
        p.error("--gsc, --ga4 and --crawl are required (or use --demo)")

    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    print("Reading exports")
    crawl, gsc, ga4 = read_crawl(a.crawl), read_gsc(a.gsc), read_ga4(a.ga4)
    prev = read_gsc(a.gsc_previous) if a.gsc_previous else pd.DataFrame(columns=GSC_COLUMNS)
    crawl.to_csv(out / "crawl.csv", index=False)
    gsc.to_csv(out / "gsc_current.csv", index=False)
    prev.to_csv(out / "gsc_previous.csv", index=False)
    ga4.to_csv(out / "ga4.csv", index=False)
    print(f"Wrote crawl.csv ({len(crawl)} rows), gsc_current.csv ({len(gsc)}), gsc_previous.csv ({len(prev)}), ga4.csv ({len(ga4)}) to {out}")

    report, missing = join_report(crawl, gsc, ga4)
    print("\nJoin check (URLs are matched on path, ignoring host, query string, trailing slash and case)")
    for k, v in report.items():
        print(f"  {k}: {v}")
    if missing:
        print(f"  {len(missing)} page(s) get clicks but are not in the crawl (orphans, or outside the crawl scope):")
        for u in missing[:10]:
            print("   -", u)

    if a.fill:
        target = out / "content-audit-filled.xlsx"
        fill_workbook(a.fill, target, {"Crawl": crawl, "GSC current": gsc, "GSC previous": prev, "GA4": ga4})
        print(f"\nFilled workbook: {target}\nImport it to Google Drive (see SETUP.md). Formulas recalculate on open.")


if __name__ == "__main__":
    main()
