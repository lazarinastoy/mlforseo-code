"""Build knowledge-panel-audit-checklist workbook (.xlsx) and checklist.md from checks.csv.

    pip install openpyxl
    python build.py

Part of https://github.com/lazarinastoy/mlforseo-code (MIT). Lazarina Stoy, MLforSEO.
"""
# ---------------------------------------------------------------------------
# Checklist workbook engine (shared by the MLforSEO checklist folders).
# Builds a Google-Sheets-ready .xlsx with openpyxl and a plain checklist.md
# from checks.csv. Only uses functions that work in both Google Sheets and
# Excel (COUNTIF, COUNTIFS, COUNTA, IF, IFERROR, ROUND, REPT, HYPERLINK, MID,
# FIND, LEN, SUBSTITUTE, RIGHT, AND, OR).
# ---------------------------------------------------------------------------
import csv
import math
import pathlib

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = pathlib.Path(__file__).resolve().parent

# MLforSEO palette, used sparingly
C = {
    "header": "251046", "ink": "1A0B32", "accent": "8448EF", "band": "F4F0FD",
    "high": "FFDE59", "white": "FFFFFF", "muted": "6B6480", "line": "DCD3F0",
    "done_fill": "DDF3E4", "done_font": "1D6B3A", "prog_fill": "E6DCFB",
    "prog_font": "4B2A99", "na_fill": "EEEEEE", "na_font": "8A8A8A",
    "warn_fill": "FDE2E2", "warn_font": "9B1C1C",
}
FONT = "Arial"  # Google Sheets' default; renders the same in Excel
STATUSES = ["Not started", "In progress", "Done", "N/A"]
PRIORITIES = ["High", "Medium", "Low"]
EVIDENCE = [
    ("Google-documented", "Stated in Google Search Central or Google Help documentation."),
    ("Platform-documented", "Stated in another platform's own documentation (OpenAI, Anthropic, Perplexity, Microsoft Bing, Apple, Meta, Wikidata / Wikimedia, schema.org)."),
    ("Research-backed", "Supported by published research or a large public dataset. Results come from specific test conditions and may not carry over to every site."),
    ("MLforSEO practice", "What we teach in MLforSEO courses and use in audits, based on how retrieval and entity systems work. Sensible and low-risk, but not confirmed by the platforms."),
    ("Unconfirmed (optional)", "A cheap extra that no major platform has confirmed using. Do it last, if at all."),
]

THIN = Side(style="thin", color=C["line"])
BORDER = Border(bottom=THIN)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
CENTER_TOP = Alignment(horizontal="center", vertical="top", wrap_text=True)


def font(size=10, bold=False, color=None, italic=False, underline=None):
    return Font(name=FONT, size=size, bold=bold, color=color or C["ink"], italic=italic, underline=underline)


def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


def q(text):
    """Escape a string for use inside a spreadsheet formula."""
    return str(text).replace('"', '""')


def link_formula(url, label):
    if not url:
        return ""
    return f'=HYPERLINK("{q(url)}","{q(label or url)}")'


def est_lines(text, width_chars):
    if not text:
        return 1
    lines = 0
    for para in str(text).split("\n"):
        lines += max(1, math.ceil(len(para) * 1.08 / max(width_chars, 1)))
    return lines


def read_checks(path=None):
    path = path or HERE / "checks.csv"
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


# --------------------------------------------------------------------------
# Generic table helpers
# --------------------------------------------------------------------------
def header_row(ws, headers, row=1, widths=None, height=30):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = font(10, True, C["white"])
        c.fill = fill(C["header"])
        c.alignment = Alignment(wrap_text=True, vertical="center")
        if widths:
            ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.row_dimensions[row].height = height


def list_validation(ws, values, rng, prompt=None):
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True,
                        showDropDown=False)  # showDropDown=False means the in-cell arrow IS shown
    dv.error = "Pick a value from the list."
    dv.errorTitle = "Not in list"
    if prompt:
        dv.prompt = prompt
    ws.add_data_validation(dv)
    dv.add(rng)
    return dv


def status_cf(ws, rng):
    rules = [
        ("Done", C["done_fill"], C["done_font"]),
        ("In progress", C["prog_fill"], C["prog_font"]),
        ("N/A", C["na_fill"], C["na_font"]),
        ("Not started", None, C["muted"]),
    ]
    for value, f_, col in rules:
        kw = {"font": Font(name=FONT, color=col, bold=value != "Not started")}
        if f_:
            kw["fill"] = fill(f_)
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{value}"'], stopIfTrue=True, **kw))


def priority_cf(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"High"'], stopIfTrue=True,
                                                  fill=fill(C["high"]), font=Font(name=FONT, bold=True, color=C["header"])))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Low"'], stopIfTrue=True,
                                                  font=Font(name=FONT, color=C["muted"])))


def banding_cf(ws, rng):
    # Added last so status / priority colours win. Survives sorting and filtering.
    ws.conditional_formatting.add(rng, FormulaRule(formula=["MOD(ROW(),2)=0"], fill=fill(C["band"])))


def print_setup(ws, landscape=True):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.gridLines = False


# --------------------------------------------------------------------------
# Checklist tab
# --------------------------------------------------------------------------
CHECK_COLS = [
    ("ID", 7), (None, 20), ("Check", 44), ("Why it matters", 46), ("How to check (tool / where)", 46),
    ("Priority", 10), ("Evidence", 17), ("Status", 13), ("Owner", 13), ("Notes", 28),
    ("Learn more", 30), ("Official source", 30),
]


def build_checklist(wb, checks, title="Checklist", area_label="Area", spare_rows=15, tab_color=None):
    ws = wb.create_sheet(title)
    cols = [(area_label if h is None else h, w) for h, w in CHECK_COLS]
    header_row(ws, [h for h, _ in cols], widths=[w for _, w in cols])
    r = 1
    for ck in checks:
        r += 1
        values = [ck["id"], ck["area"], ck["check"], ck["why"], ck["how"], ck["priority"], ck["evidence"],
                  "Not started", "", "", link_formula(ck.get("learn_url"), ck.get("learn_label")),
                  link_formula(ck.get("source_url"), ck.get("source_label"))]
        for ci, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=ci, value=v)
            cell.font = font(10, bold=(ci == 3))
            cell.alignment = CENTER_TOP if ci in (1, 6, 8) else WRAP_TOP
            cell.border = BORDER
            if ci in (11, 12) and v:
                cell.font = font(10, color=C["accent"], underline="single")
        lines = max(est_lines(ck["check"], 44), est_lines(ck["why"], 46), est_lines(ck["how"], 46))
        ws.row_dimensions[r].height = min(15 * lines + 6, 300)
    last_check = r
    for _ in range(spare_rows):  # blank rows for the user's own checks
        r += 1
        for ci in range(1, len(cols) + 1):
            cell = ws.cell(row=r, column=ci, value=None)
            cell.font = font(10)
            cell.alignment = WRAP_TOP
            cell.border = BORDER
    last = r
    ncol = get_column_letter(len(cols))
    list_validation(ws, STATUSES, f"H2:H{last}")
    list_validation(ws, PRIORITIES, f"F2:F{last}")
    list_validation(ws, [e for e, _ in EVIDENCE], f"G2:G{last}")
    status_cf(ws, f"H2:H{last}")
    priority_cf(ws, f"F2:F{last}")
    banding_cf(ws, f"A2:{ncol}{last}")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{ncol}{last}"
    if tab_color:
        ws.sheet_properties.tabColor = tab_color
    print_setup(ws)
    return ws, last_check, last


# --------------------------------------------------------------------------
# Start here tab
# --------------------------------------------------------------------------
def build_start(wb, blocks, tab_color=C["accent"]):
    """blocks: list of tuples. Types: title, subtitle, h, p, steps, kv (list of (k, v)),
    status_legend, priority_legend, evidence_legend, links (list of (label, url)), spacer."""
    ws = wb.active
    ws.title = "Start here"
    ws.sheet_properties.tabColor = tab_color
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 92
    r = 1

    def para(text, size=10, bold=False, color=None, italic=False, height=None):
        nonlocal r
        r += 1
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        c = ws.cell(row=r, column=2, value=text)
        c.font = font(size, bold, color, italic)
        c.alignment = WRAP_TOP
        ws.row_dimensions[r].height = height or (max(1, est_lines(text, 112)) * (size + 5) + 4)

    def kv(k, v, kfill=None, kfont=None, link=None):
        nonlocal r
        r += 1
        a = ws.cell(row=r, column=2, value=k)
        a.font = kfont or font(10, True)
        a.alignment = WRAP_TOP
        if kfill:
            a.fill = fill(kfill)
        b = ws.cell(row=r, column=3, value=link_formula(link, v) if link else v)
        b.font = font(10, color=C["accent"], underline="single") if link else font(10)
        b.alignment = WRAP_TOP
        ws.row_dimensions[r].height = max(est_lines(k, 26), est_lines(v, 90)) * 15 + 4

    for blk in blocks:
        kind = blk[0]
        if kind == "title":
            para(blk[1], 18, True, C["header"], height=30)
        elif kind == "subtitle":
            para(blk[1], 11, False, C["muted"])
        elif kind == "h":
            r += 1
            para(blk[1], 12, True, C["accent"], height=22)
        elif kind == "p":
            para(blk[1])
        elif kind == "note":
            para(blk[1], 9, False, C["muted"], italic=True)
        elif kind == "steps":
            for i, s in enumerate(blk[1], start=1):
                kv(f"Step {i}", s)
        elif kind == "kv":
            for k, v in blk[1]:
                kv(k, v)
        elif kind == "status_legend":
            kv("Not started", "Default for every check.", kfont=Font(name=FONT, color=C["muted"]))
            kv("In progress", "Someone is working on it.", kfill=C["prog_fill"], kfont=Font(name=FONT, bold=True, color=C["prog_font"]))
            kv("Done", "Checked, and fixed if it needed fixing.", kfill=C["done_fill"], kfont=Font(name=FONT, bold=True, color=C["done_font"]))
            kv("N/A", "Does not apply to this site. N/A rows are left out of the % complete.", kfill=C["na_fill"], kfont=Font(name=FONT, color=C["na_font"]))
        elif kind == "priority_legend":
            kv("High", blk[1][0], kfill=C["high"], kfont=Font(name=FONT, bold=True, color=C["header"]))
            kv("Medium", blk[1][1])
            kv("Low", blk[1][2], kfont=Font(name=FONT, color=C["muted"]))
        elif kind == "evidence_legend":
            for k, v in EVIDENCE:
                kv(k, v)
        elif kind == "links":
            for label, url in blk[1]:
                kv(label[0], label[1], link=url)
        elif kind == "spacer":
            r += 1
    print_setup(ws, landscape=False)
    return ws


# --------------------------------------------------------------------------
# Progress tab (COUNTIF / COUNTIFS only, so it works in Sheets and Excel)
# --------------------------------------------------------------------------
def build_progress(wb, sheet, areas, area_label="Area", last_row=500, title="Progress"):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = C["high"]
    s = f"'{sheet}'"
    A, P, S, ID = f"{s}!$B$2:$B${last_row}", f"{s}!$F$2:$F${last_row}", f"{s}!$H$2:$H${last_row}", f"{s}!$C$2:$C${last_row}"
    widths = [3, 30, 10, 10, 13, 13, 10, 12, 28, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws["B2"] = "Progress"
    ws["B2"].font = font(18, True, C["header"])
    ws["B3"] = "Updates automatically as you change the Status column. Formulas use COUNTIF / COUNTIFS, so they work in Google Sheets and Excel."
    ws["B3"].font = font(9, color=C["muted"], italic=True)

    # Overall
    labels = [("Checks", f"=COUNTA({ID})"), ("Done", f'=COUNTIF({S},"Done")'),
              ("In progress", f'=COUNTIF({S},"In progress")'), ("Not started", f'=COUNTIF({S},"Not started")'),
              ("N/A", f'=COUNTIF({S},"N/A")')]
    r = 5
    for k, f_ in labels:
        ws.cell(row=r, column=2, value=k).font = font(10, True)
        c = ws.cell(row=r, column=3, value=f_)
        c.font = font(10)
        c.alignment = Alignment(horizontal="right")
        r += 1
    ws.cell(row=r, column=2, value="% complete (Done / applicable)").font = font(10, True)
    pc = ws.cell(row=r, column=3, value="=IFERROR(C6/(C5-C9),0)")
    pc.number_format = "0%"
    pc.font = font(12, True, C["accent"])
    bar = ws.cell(row=r, column=4, value=f'=REPT("█",ROUND(C{r}*20,0))&REPT("░",20-ROUND(C{r}*20,0))')
    bar.font = Font(name=FONT, size=11, color=C["accent"])
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=8)
    pct_row = r
    r += 1
    ws.cell(row=r, column=2, value="Open High-priority checks").font = font(10, True)
    c = ws.cell(row=r, column=3, value=f'=COUNTIFS({P},"High",{S},"<>Done",{S},"<>N/A")')
    c.font = font(12, True, C["header"])
    c.fill = fill(C["high"])

    # By area
    r += 2
    hdr = [area_label, "Checks", "Done", "In progress", "Not started", "N/A", "% done", "Open High", ]
    for i, h in enumerate(hdr, start=2):
        c = ws.cell(row=r, column=i, value=h)
        c.font = font(10, True, C["white"])
        c.fill = fill(C["header"])
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[r].height = 28
    first = r + 1
    for a in areas:
        r += 1
        ws.cell(row=r, column=2, value=a).font = font(10, True)
        ref = f"$B{r}"
        fs = [f"=COUNTIF({A},{ref})",
              f'=COUNTIFS({A},{ref},{S},"Done")', f'=COUNTIFS({A},{ref},{S},"In progress")',
              f'=COUNTIFS({A},{ref},{S},"Not started")', f'=COUNTIFS({A},{ref},{S},"N/A")',
              f"=IFERROR(D{r}/(C{r}-G{r}),0)",
              f'=COUNTIFS({A},{ref},{P},"High",{S},"<>Done",{S},"<>N/A")']
        for i, f_ in enumerate(fs, start=3):
            c = ws.cell(row=r, column=i, value=f_)
            c.font = font(10)
            c.alignment = Alignment(horizontal="right")
            if i == 8:
                c.number_format = "0%"
        bar = ws.cell(row=r, column=10, value=f'=REPT("█",ROUND(H{r}*10,0))&REPT("░",10-ROUND(H{r}*10,0))')
        bar.font = Font(name=FONT, size=10, color=C["accent"])
    ws.column_dimensions["J"].width = 16
    ws.conditional_formatting.add(f"B{first}:J{r}", FormulaRule(formula=["MOD(ROW(),2)=0"], fill=fill(C["band"])))
    ws.conditional_formatting.add(f"I{first}:I{r}", CellIsRule(operator="greaterThan", formula=["0"],
                                                               fill=fill(C["high"]), font=Font(name=FONT, bold=True, color=C["header"])))

    # By priority
    r += 2
    for i, h in enumerate(["Priority", "Checks", "Done", "Open", "% done"], start=2):
        c = ws.cell(row=r, column=i, value=h)
        c.font = font(10, True, C["white"])
        c.fill = fill(C["header"])
    for p in PRIORITIES:
        r += 1
        ws.cell(row=r, column=2, value=p).font = font(10, True)
        ref = f"$B{r}"
        fs = [f"=COUNTIF({P},{ref})", f'=COUNTIFS({P},{ref},{S},"Done")',
              f'=COUNTIFS({P},{ref},{S},"<>Done",{S},"<>N/A")',
              f'=IFERROR(D{r}/(C{r}-COUNTIFS({P},{ref},{S},"N/A")),0)']
        for i, f_ in enumerate(fs, start=3):
            c = ws.cell(row=r, column=i, value=f_)
            c.font = font(10)
            c.alignment = Alignment(horizontal="right")
            if i == 6:
                c.number_format = "0%"
    priority_cf(ws, f"B{r-2}:B{r}")
    print_setup(ws, landscape=False)
    return ws, r


# --------------------------------------------------------------------------
# Simple log / tracker tabs
# --------------------------------------------------------------------------
def build_log(wb, title, columns, rows=200, validations=None, intro=None, example_rows=None, tab_color=None,
              formulas=None, wrap_cols=None):
    """columns: list of (header, width). validations: {col_letter: [values]}.
    formulas: {col_letter: template with {r}} filled for every row."""
    ws = wb.create_sheet(title)
    start = 1
    if intro:
        ws.cell(row=1, column=1, value=intro).font = font(9, color=C["muted"], italic=True)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
        ws.row_dimensions[1].height = max(1, est_lines(intro, sum(w for _, w in columns))) * 14 + 6
        ws.cell(row=1, column=1).alignment = WRAP_TOP
        start = 2
    header_row(ws, [h for h, _ in columns], row=start, widths=[w for _, w in columns])
    first = start + 1
    last = start + rows
    example_rows = example_rows or []
    for i in range(rows):
        r = first + i
        ex = example_rows[i] if i < len(example_rows) else None
        for ci in range(1, len(columns) + 1):
            col = get_column_letter(ci)
            v = None
            if formulas and col in formulas:
                v = formulas[col].format(r=r)
            elif ex is not None and ci - 1 < len(ex):
                v = ex[ci - 1]
            cell = ws.cell(row=r, column=ci, value=v if v != "" else None)
            cell.font = font(10, italic=ex is not None and not (formulas and col in formulas),
                             color=C["muted"] if ex is not None else None)
            cell.alignment = WRAP_TOP if (wrap_cols is None or col in wrap_cols) else Alignment(vertical="top")
            cell.border = BORDER
    for col, values in (validations or {}).items():
        list_validation(ws, values, f"{col}{first}:{col}{last}")
    ncol = get_column_letter(len(columns))
    ws.freeze_panes = ws.cell(row=first, column=2)
    ws.auto_filter.ref = f"A{start}:{ncol}{last}"
    if tab_color:
        ws.sheet_properties.tabColor = tab_color
    print_setup(ws)
    return ws, first, last


# --------------------------------------------------------------------------
# Markdown version
# --------------------------------------------------------------------------
def md_link(label, url):
    return f"[{label}]({url})" if url else ""


def write_markdown(path, title, intro_md, checks, area_order, outro_md, area_intros=None):
    out = [f"# {title}", "", intro_md.strip(), ""]
    for area in area_order:
        out.append(f"## {area}")
        out.append("")
        if area_intros and area in area_intros:
            out += [area_intros[area], ""]
        for ck in [c for c in checks if c["area"] == area]:
            out.append(f"- [ ] **{ck['id']} · {ck['check']}**  ")
            out.append(f"  *Priority: {ck['priority']} · Evidence: {ck['evidence']}*  ")
            out.append(f"  **Why it matters:** {ck['why']}  ")
            out.append(f"  **How to check:** {ck['how']}  ")
            links = [x for x in (md_link(ck.get('learn_label'), ck.get('learn_url')),
                                 md_link(ck.get('source_label'), ck.get('source_url'))) if x]
            if links:
                out.append(f"  **Learn more:** {' · '.join(links)}")
            out.append("")
    out += [outro_md.strip(), ""]
    pathlib.Path(path).write_text("\n".join(out), encoding="utf-8")


# ===========================================================================
# Knowledge panel & brand SERP audit: folder-specific content and tabs
# Run:  python build.py   (needs openpyxl)  ->  knowledge-panel-audit-checklist.xlsx + checklist.md
# ===========================================================================
TITLE = "Knowledge Panel & Brand SERP Audit Checklist"
XLSX = "knowledge-panel-audit-checklist.xlsx"
AREAS = ['1. Brand SERP baseline', '2. Knowledge Graph status', '3. Wikidata & Wikipedia', '4. On-site entity markup', '5. Cross-platform consistency', '6. Corroboration', '7. Claim & manage', '8. Monitor']
ACTIONS = [
    ('1', 'Agree the canonical facts and fill in the Fact sheet', 'K22', 'Everything else is checked against it.', 'Low'),
    ('2', 'Fix contradictions on your own site and markup (about page, Organization markup, @id, sameAs)', 'K16 to K21, K24', 'You control these, and Google re-reads them at the next crawl.', 'Low to medium'),
    ('3', 'Align the profiles you control (LinkedIn, Crunchbase, social profiles, directories, Business Profile)', 'K23 to K25, K31', 'Removes conflicts across the verification points you can edit today.', 'Medium'),
    ('4', 'Claim the knowledge panel if one exists and submit evidence-backed corrections', 'K29, K30, K32', 'Verified feedback is prioritised.', 'Low'),
    ('5', 'Improve or create the Wikidata item if you meet notability, declaring any conflict of interest', 'K11 to K14', 'Adds a referenced, openly reused record of your entity.', 'Medium'),
    ('6', 'Earn independent coverage that confirms your key facts', 'K26 to K28', 'Panels and AI citations rely on agreement across independent sources.', 'High (ongoing)'),
    ('7', 'Fill the gaps on page one for brand and brand + modifier searches', 'K33', 'Keeps accurate, owned information in front of people.', 'Medium'),
    ('8', 'Set the monthly and quarterly checks in your calendar', 'K34 to K36', 'Stops drift and shows whether the work is moving the entity.', 'Low'),
]
FACTS = [
    'Canonical brand name',
    'Legal name',
    'Other names / abbreviations',
    'Website URL (exact)',
    'One-sentence description',
    'Category / industry',
    'Founding date',
    'Founders',
    'Current CEO / leader',
    'Headquarters address',
    'Phone',
    'Employee range',
    'Logo URL',
    'Wikidata Q-ID',
    'Knowledge Graph ID',
]
PLATFORMS = [
    'Website & schema',
    'LinkedIn',
    'Crunchbase',
    'Wikidata',
    'Business Profile',
    'X / social',
    'YouTube',
    'Industry directory',
    'Company registry',
]
RESULT_TYPES = ["Owned site", "Official profile", "Third party: positive", "Third party: neutral", "Third party: negative", "Namesake"]
FEATURES = ["Organic result", "Knowledge panel", "Business Profile", "Sitelinks", "AI Overview", "People Also Ask",
            "Top stories", "Videos", "Images", "Social posts", "Reviews", "Other"]

START = [
    ("title", TITLE),
    ("subtitle", "A step-by-step audit of your Google knowledge panel eligibility and brand SERP, ending with a prioritised action plan. Free from MLforSEO, October 2026."),
    ("h", "What this is"),
    ("p", "Knowledge panels aren't something you apply for. Google generates them automatically when its Knowledge Graph holds enough consistent, verifiable information about an entity, "
          "and the same signals help AI assistants describe you correctly. This workbook walks you through 36 checks: what page one shows for your brand today, "
          "whether Google's Knowledge Graph recognises you, your Wikidata status, your own markup, cross-platform consistency, independent corroboration, claiming and managing a panel, and monitoring."),
    ("p", "It works for organisations and for people (founders, authors, consultants). Where a check says Organization, read Person or ProfilePage if you're auditing a person."),
    ("h", "How to use it"),
    ("steps", [
        "Make your own copy (File > Make a copy in Google Sheets).",
        "Start with the Fact sheet tab: write the canonical version of each fact in column B. Everything else is compared against it.",
        "Work through the Checklist tab in order. Log what you find in the Brand SERP log and KG lookup log tabs as the checks ask.",
        "Fill in the platform columns of the Fact sheet. Any cell that differs from column B turns red.",
        "Turn findings into the Action plan tab. Its order (own site first, then profiles, then panel claim, Wikidata and coverage) is the sequence we recommend.",
        "Check the Progress tab for completion, owned share of page one and fact mismatches. Repeat monthly and quarterly (area 8).",
    ]),
    ("h", "Tabs"),
    ("kv", [
        ("Checklist", "36 checks with why each matters, how to check it, priority, evidence level, status, owner, notes and links."),
        ("Fact sheet", "Your single source of truth, compared across nine platforms. Mismatches turn red and are counted."),
        ("Brand SERP log", "One row per result on page one for each brand query you test."),
        ("KG lookup log", "Dated Knowledge Graph Search API results: ID, name, type, description, resultScore and panel status."),
        ("Action plan", "The prioritised sequence of fixes, with owners, dates and status."),
        ("Progress", "Checklist completion by area and priority, plus brand SERP and fact-sheet summaries."),
    ]),
    ("h", "Status"),
    ("status_legend",),
    ("h", "Priority"),
    ("priority_legend", ["Fix first. Contradictions, misidentification or missing foundations.",
                         "Do next. Strengthens recognition and corroboration.",
                         "Monitor, or nice to have."]),
    ("h", "Evidence levels"),
    ("evidence_legend",),
    ("h", "Realistic expectations"),
    ("p", "Google doesn't publish a threshold for showing a panel, and it can't be bought or requested. In MLforSEO's experience it takes six to eighteen months of consistent signals for most brands, faster if you already have strong press coverage. "
          "The same work makes you easier for AI assistants to recognise and describe accurately, so it pays off even before a panel appears."),
    ("h", "Credits"),
    ("links", [
        (("Created by", "Lazarina Stoy · MLforSEO"), "https://www.mlforseo.com/experts/lazarina-stoy/"),
        (("Built on", "Beatrice Gamba's entity SEO articles on MLforSEO"), "https://www.mlforseo.com/experts/beatrice-gamba/"),
        (("Read first", "What it takes to earn a Google knowledge panel"), "https://www.mlforseo.com/knowledge-graph-strategy/what-it-takes-to-earn-a-google-knowledge-panel/"),
        (("Go deeper", "AI Search & LLMs: Entity SEO and Knowledge Graph Strategies for Brands"), "https://www.mlforseo.com/courses/ai-search-llms-entity-seo-and-knowledge-graph-strategies-for-brands/"),
        (("Source files", "GitHub: lazarinastoy/mlforseo-code"), "https://github.com/lazarinastoy/mlforseo-code/tree/main/knowledge-panel-audit-checklist"),
    ]),
    ("note", "Version 1.0, October 2026. Free to use and adapt (MIT licence); please keep the credit."),
]


def build_facts(wb):
    ws = wb.create_sheet("Fact sheet")
    cols = [("Fact", 26), ("Canonical value (your source of truth)", 32)] + [(p, 20) for p in PLATFORMS] + [("Mismatches", 12), ("Notes", 28)]
    ncols = len(cols)
    intro = ("Write the canonical value in column B, then what each platform currently shows. Leave a cell blank if that platform doesn't show the fact. "
             "Cells that differ from column B turn red (the comparison ignores upper and lower case). Add rows for any other facts you track.")
    ws.cell(row=1, column=1, value=intro).font = font(9, color=C["muted"], italic=True)
    ws.cell(row=1, column=1).alignment = WRAP_TOP
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws.row_dimensions[1].height = 32
    header_row(ws, [h for h, _ in cols], row=2, widths=[w for _, w in cols], height=32)
    first = 3
    rows = FACTS + [""] * 5
    plat_first, plat_last = get_column_letter(3), get_column_letter(2 + len(PLATFORMS))
    mm_col = get_column_letter(3 + len(PLATFORMS))
    for i, fact in enumerate(rows):
        r = first + i
        for ci in range(1, ncols + 1):
            c = ws.cell(row=r, column=ci, value=None)
            c.font = font(10, bold=ci == 1)
            c.alignment = WRAP_TOP
            c.border = BORDER
        ws.cell(row=r, column=1, value=fact or None)
        ws.cell(row=r, column=3 + len(PLATFORMS),
                value=f'=IF($B{r}="","",SUMPRODUCT(({plat_first}{r}:{plat_last}{r}<>"")*({plat_first}{r}:{plat_last}{r}<>$B{r})))').alignment = CENTER_TOP
        ws.row_dimensions[r].height = 30
    last = first + len(rows) - 1
    ws.conditional_formatting.add(f"{plat_first}{first}:{plat_last}{last}",
                                  FormulaRule(formula=[f'AND({plat_first}{first}<>"",$B{first}<>"",{plat_first}{first}<>$B{first})'], stopIfTrue=True,
                                              fill=fill(C["warn_fill"]), font=Font(name=FONT, color=C["warn_font"], bold=True)))
    ws.conditional_formatting.add(f"{mm_col}{first}:{mm_col}{last}", CellIsRule(operator="greaterThan", formula=["0"],
                                  fill=fill(C["high"]), font=Font(name=FONT, bold=True, color=C["header"])))
    ws.conditional_formatting.add(f"B{first}:B{last}", FormulaRule(formula=["TRUE"], fill=fill(C["band"])))
    ws.freeze_panes = "C3"
    ws.sheet_properties.tabColor = C["accent"]
    print_setup(ws)
    return ws, first, last, mm_col


def build_serp_log(wb):
    cols = [("Date", 11), ("Query", 24), ("Device", 10), ("Position", 9), ("URL", 36), ("Title / source", 30),
            ("Result type", 20), ("SERP feature", 18), ("Accurate?", 10), ("Notes", 32)]
    ex = [["2026-10-01", "acme software", "Desktop", 1, "https://www.example.com/", "Acme Software: project management for remote teams", "Owned site", "Sitelinks", "Yes", "EXAMPLE ROW: delete."],
          ["2026-10-01", "acme software", "Desktop", 2, "https://www.linkedin.com/company/example", "Acme Software | LinkedIn", "Official profile", "Organic result", "Partly", "EXAMPLE ROW: delete. Old HQ address."]]
    ws, first, last = build_log(wb, "Brand SERP log", cols, rows=300,
                                intro="One row per result on page one, per query and device. Log the exact brand name first, then brand + modifier queries (reviews, pricing, vs competitor, founder name). Delete the example rows.",
                                validations={"C": ["Desktop", "Mobile"], "G": RESULT_TYPES, "H": FEATURES, "I": ["Yes", "Partly", "No"]},
                                example_rows=ex, tab_color=C["accent"])
    ws.conditional_formatting.add(f"G{first}:G{last}", CellIsRule(operator="equal", formula=['"Third party: negative"'], fill=fill(C["warn_fill"]), font=Font(name=FONT, color=C["warn_font"], bold=True)))
    ws.conditional_formatting.add(f"I{first}:I{last}", CellIsRule(operator="equal", formula=['"No"'], fill=fill(C["warn_fill"]), font=Font(name=FONT, color=C["warn_font"], bold=True)))
    banding_cf(ws, f"A{first}:J{last}")
    return ws, first, last


def build_kg_log(wb):
    cols = [("Date", 11), ("Query", 22), ("KG ID (kg:/m/... or kg:/g/...)", 22), ("Name", 22), ("@type", 18),
            ("Description", 28), ("resultScore", 11), ("Is it you?", 10), ("Panel for brand query", 18), ("Claimed?", 13), ("Notes", 30)]
    ws, first, last = build_log(wb, "KG lookup log", cols, rows=120,
                                intro="Log a Knowledge Graph Search API lookup for your brand, its name variants and founders each month. resultScore has no official thresholds: compare it with your own earlier rows only.",
                                validations={"H": ["Yes", "No", "Unsure"], "I": ["Knowledge panel", "Business Profile", "None"], "J": ["Yes", "No", "Not claimable yet"]},
                                tab_color=C["accent"])
    ws.conditional_formatting.add(f"H{first}:H{last}", CellIsRule(operator="equal", formula=['"No"'], fill=fill(C["warn_fill"]), font=Font(name=FONT, color=C["warn_font"], bold=True)))
    banding_cf(ws, f"A{first}:K{last}")
    return ws


def build_actions(wb):
    cols = [("Order", 7), ("Action", 44), ("Related checks", 16), ("Why this order", 40), ("Effort", 13),
            ("Owner", 14), ("Due date", 12), ("Status", 13), ("Notes", 30)]
    rows = [list(a) + ["", "", "Not started", ""] for a in ACTIONS]
    ws, first, last = build_log(wb, "Action plan", cols, rows=len(rows) + 12, example_rows=None,
                                intro="The recommended sequence: fix what you control first, then align profiles, then claim, then build external corroboration. Add your own actions from the checklist findings in the blank rows.",
                                validations={"H": STATUSES}, tab_color=C["high"])
    for i, vals in enumerate(rows):
        r = first + i
        for ci, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=ci, value=v if v != "" else None)
            c.font = font(10, bold=ci == 2)
            c.alignment = CENTER_TOP if ci in (1, 8) else WRAP_TOP
        ws.row_dimensions[r].height = max(est_lines(vals[1], 44), est_lines(vals[3], 40)) * 15 + 6
    status_cf(ws, f"H{first}:H{last}")
    banding_cf(ws, f"A{first}:I{last}")
    return ws, first, last


def add_kp_summary(ws, r, serp, facts, actions):
    s_first, s_last = serp
    f_first, f_last, mm_col = facts
    a_first, a_last = actions
    S = "'Brand SERP log'"
    T = f"{S}!$G${s_first}:$G${s_last}"
    ACC = f"{S}!$I${s_first}:$I${s_last}"
    r += 2
    ws.cell(row=r, column=2, value="Brand SERP, facts and actions").font = font(12, True, C["accent"])
    rows = [
        ("Results logged", f"=COUNTA({T})", None),
        ("Owned site or official profile", f'=COUNTIF({T},"Owned site")+COUNTIF({T},"Official profile")', None),
        ("Owned share of logged results", f"=IFERROR(C{r+2}/C{r+1},0)", "0%"),
        ("Negative third-party results", f'=COUNTIF({T},"Third party: negative")', None),
        ("Results with inaccurate information", f'=COUNTIF({ACC},"No")+COUNTIF({ACC},"Partly")', None),
        ("Fact-sheet mismatches (all platforms)", f"=SUM('Fact sheet'!{mm_col}{f_first}:{mm_col}{f_last})", None),
        ("Facts with at least one mismatch", f"=COUNTIF('Fact sheet'!{mm_col}{f_first}:{mm_col}{f_last},\">0\")", None),
        ("Action plan items done", f"=COUNTIF('Action plan'!$H${a_first}:$H${a_last},\"Done\")&\" of \"&COUNTA('Action plan'!$B${a_first}:$B${a_last})", None),
    ]
    for label, f_, fmt in rows:
        r += 1
        ws.cell(row=r, column=2, value=label).font = font(10, True)
        c = ws.cell(row=r, column=3, value=f_)
        c.font = font(10)
        c.alignment = Alignment(horizontal="right")
        if fmt:
            c.number_format = fmt
    ws.conditional_formatting.add(f"C{r-4}:C{r-1}", CellIsRule(operator="greaterThan", formula=["0"], fill=fill(C["warn_fill"]), font=Font(name=FONT, color=C["warn_font"], bold=True)))
    return r


def main():
    checks = read_checks()
    wb = Workbook()
    build_start(wb, START)
    ws, last_check, last = build_checklist(wb, checks, area_label="Area")
    facts = build_facts(wb)
    _, s_first, s_last = build_serp_log(wb)
    build_kg_log(wb)
    _, a_first, a_last = build_actions(wb)
    prog, r = build_progress(wb, "Checklist", AREAS, last_row=last)
    add_kp_summary(prog, r, (s_first, s_last), facts[1:], (a_first, a_last))
    wb.save(HERE / XLSX)
    write_markdown(HERE / "checklist.md", TITLE, MD_INTRO, checks, AREAS, MD_OUTRO)
    print("wrote", XLSX, "and checklist.md with", len(checks), "checks")


MD_INTRO = """
*36 checks for your Google knowledge panel eligibility and brand SERP, ending with a prioritised action plan. By [Lazarina Stoy](https://www.mlforseo.com/experts/lazarina-stoy/) for [MLforSEO](https://www.mlforseo.com/), built on [Beatrice Gamba](https://www.mlforseo.com/experts/beatrice-gamba/)'s entity SEO articles. Checked October 2026.*

This is the plain-text version of `knowledge-panel-audit-checklist.xlsx`. The Google Sheets version adds a fact sheet that flags mismatches across platforms, a brand SERP log, a Knowledge Graph lookup log, an action plan and a progress dashboard.

Knowledge panels are generated automatically when Google's Knowledge Graph holds enough consistent, verifiable information about an entity; you can't apply for one. The work below builds those signals, and the same signals help AI assistants describe you accurately. Each check has an **evidence level**:

| Evidence | Meaning |
|---|---|
| Google-documented | Stated in Google Search Central or Google Help documentation |
| Platform-documented | Stated in another platform's own docs (Wikidata / Wikimedia, schema.org) |
| MLforSEO practice | What we teach and use in audits; sensible and low-risk, but not confirmed by the platforms |

**Priority:** High = fix first. Medium = do next. Low = monitor or nice to have.
"""

MD_OUTRO = """
---

## Prioritised action plan

1. Agree the canonical facts and fill in the fact sheet (K22).
2. Fix contradictions on your own site and markup: about page, Organization markup, @id, sameAs (K16 to K21, K24).
3. Align the profiles you control: LinkedIn, Crunchbase, social profiles, directories, Business Profile (K23 to K25, K31).
4. Claim the knowledge panel if one exists and submit evidence-backed corrections (K29, K30, K32).
5. Improve or create the Wikidata item if you meet notability, declaring any conflict of interest (K11 to K14).
6. Earn independent coverage that confirms your key facts (K26 to K28).
7. Fill the gaps on page one for brand and brand + modifier searches (K33).
8. Put the monthly and quarterly checks in your calendar (K34 to K36).

## Realistic expectations

Google doesn't publish a threshold for showing a panel. In MLforSEO's experience it takes six to eighteen months of consistent signals for most brands. The same work helps AI assistants recognise and describe you, so it pays off before a panel appears.

## Go deeper

- [What it takes to earn a Google knowledge panel](https://www.mlforseo.com/knowledge-graph-strategy/what-it-takes-to-earn-a-google-knowledge-panel/) (Beatrice Gamba)
- [AI Search & LLMs: Entity SEO and Knowledge Graph Strategies for Brands](https://www.mlforseo.com/courses/ai-search-llms-entity-seo-and-knowledge-graph-strategies-for-brands/) (MLforSEO Academy)
- [Knowledge Graph API notebook](https://www.mlforseo.com/resource/knowledge-graph-api-extraction-and-entity-relationship-analysis-for-semantic-keyword-research-notebook/)

Created by Lazarina Stoy · MLforSEO. Free to use and adapt under the repository's MIT licence; please keep the credit.
"""

if __name__ == "__main__":
    main()
