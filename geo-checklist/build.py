"""Build geo-checklist workbook (.xlsx) and checklist.md from checks.csv.

    pip install openpyxl
    python build.py

Part of https://github.com/lazarinastoy/mlforseo-code (MIT). Lazarina Stoy, MLforSEO.
"""
# ---------------------------------------------------------------------------
# Checklist workbook engine (shared by the MLforSEO checklist folders).
# Builds a Google-Sheets-ready .xlsx with openpyxl and a plain checklist.md
# from checks.csv. Only uses functions that work in both Google Sheets and
# Excel (COUNTIF, COUNTIFS, COUNTA, IF, IFERROR, ROUND, REPT, HYPERLINK, MID,
# FIND, LEN, SUBSTITUTE, RIGHT, AND, OR).
# ---------------------------------------------------------------------------
import csv
import math
import pathlib

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = pathlib.Path(__file__).resolve().parent

# MLforSEO palette, used sparingly
C = {
    "header": "251046", "ink": "1A0B32", "accent": "8448EF", "band": "F4F0FD",
    "high": "FFDE59", "white": "FFFFFF", "muted": "6B6480", "line": "DCD3F0",
    "done_fill": "DDF3E4", "done_font": "1D6B3A", "prog_fill": "E6DCFB",
    "prog_font": "4B2A99", "na_fill": "EEEEEE", "na_font": "8A8A8A",
    "warn_fill": "FDE2E2", "warn_font": "9B1C1C",
}
FONT = "Arial"  # Google Sheets' default; renders the same in Excel
STATUSES = ["Not started", "In progress", "Done", "N/A"]
PRIORITIES = ["High", "Medium", "Low"]
EVIDENCE = [
    ("Google-documented", "Stated in Google Search Central or Google Help documentation."),
    ("Platform-documented", "Stated in another platform's own documentation (OpenAI, Anthropic, Perplexity, Microsoft Bing, Apple, Meta, Wikidata / Wikimedia, schema.org)."),
    ("Research-backed", "Supported by published research or a large public dataset. Results come from specific test conditions and may not carry over to every site."),
    ("MLforSEO practice", "What we teach in MLforSEO courses and use in audits, based on how retrieval and entity systems work. Sensible and low-risk, but not confirmed by the platforms."),
    ("Unconfirmed (optional)", "A cheap extra that no major platform has confirmed using. Do it last, if at all."),
]

THIN = Side(style="thin", color=C["line"])
BORDER = Border(bottom=THIN)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
CENTER_TOP = Alignment(horizontal="center", vertical="top", wrap_text=True)


def font(size=10, bold=False, color=None, italic=False, underline=None):
    return Font(name=FONT, size=size, bold=bold, color=color or C["ink"], italic=italic, underline=underline)


def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


def q(text):
    """Escape a string for use inside a spreadsheet formula."""
    return str(text).replace('"', '""')


def link_formula(url, label):
    if not url:
        return ""
    return f'=HYPERLINK("{q(url)}","{q(label or url)}")'


def est_lines(text, width_chars):
    if not text:
        return 1
    lines = 0
    for para in str(text).split("\n"):
        lines += max(1, math.ceil(len(para) * 1.08 / max(width_chars, 1)))
    return lines


def read_checks(path=None):
    path = path or HERE / "checks.csv"
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


# --------------------------------------------------------------------------
# Generic table helpers
# --------------------------------------------------------------------------
def header_row(ws, headers, row=1, widths=None, height=30):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = font(10, True, C["white"])
        c.fill = fill(C["header"])
        c.alignment = Alignment(wrap_text=True, vertical="center")
        if widths:
            ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.row_dimensions[row].height = height


def list_validation(ws, values, rng, prompt=None):
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True,
                        showDropDown=False)  # showDropDown=False means the in-cell arrow IS shown
    dv.error = "Pick a value from the list."
    dv.errorTitle = "Not in list"
    if prompt:
        dv.prompt = prompt
    ws.add_data_validation(dv)
    dv.add(rng)
    return dv


def status_cf(ws, rng):
    rules = [
        ("Done", C["done_fill"], C["done_font"]),
        ("In progress", C["prog_fill"], C["prog_font"]),
        ("N/A", C["na_fill"], C["na_font"]),
        ("Not started", None, C["muted"]),
    ]
    for value, f_, col in rules:
        kw = {"font": Font(name=FONT, color=col, bold=value != "Not started")}
        if f_:
            kw["fill"] = fill(f_)
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{value}"'], stopIfTrue=True, **kw))


def priority_cf(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"High"'], stopIfTrue=True,
                                                  fill=fill(C["high"]), font=Font(name=FONT, bold=True, color=C["header"])))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Low"'], stopIfTrue=True,
                                                  font=Font(name=FONT, color=C["muted"])))


def banding_cf(ws, rng):
    # Added last so status / priority colours win. Survives sorting and filtering.
    ws.conditional_formatting.add(rng, FormulaRule(formula=["MOD(ROW(),2)=0"], fill=fill(C["band"])))


def print_setup(ws, landscape=True):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.gridLines = False


# --------------------------------------------------------------------------
# Checklist tab
# --------------------------------------------------------------------------
CHECK_COLS = [
    ("ID", 7), (None, 20), ("Check", 44), ("Why it matters", 46), ("How to check (tool / where)", 46),
    ("Priority", 10), ("Evidence", 17), ("Status", 13), ("Owner", 13), ("Notes", 28),
    ("Learn more", 30), ("Official source", 30),
]


def build_checklist(wb, checks, title="Checklist", area_label="Area", spare_rows=15, tab_color=None):
    ws = wb.create_sheet(title)
    cols = [(area_label if h is None else h, w) for h, w in CHECK_COLS]
    header_row(ws, [h for h, _ in cols], widths=[w for _, w in cols])
    r = 1
    for ck in checks:
        r += 1
        values = [ck["id"], ck["area"], ck["check"], ck["why"], ck["how"], ck["priority"], ck["evidence"],
                  "Not started", "", "", link_formula(ck.get("learn_url"), ck.get("learn_label")),
                  link_formula(ck.get("source_url"), ck.get("source_label"))]
        for ci, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=ci, value=v)
            cell.font = font(10, bold=(ci == 3))
            cell.alignment = CENTER_TOP if ci in (1, 6, 8) else WRAP_TOP
            cell.border = BORDER
            if ci in (11, 12) and v:
                cell.font = font(10, color=C["accent"], underline="single")
        lines = max(est_lines(ck["check"], 44), est_lines(ck["why"], 46), est_lines(ck["how"], 46))
        ws.row_dimensions[r].height = min(15 * lines + 6, 300)
    last_check = r
    for _ in range(spare_rows):  # blank rows for the user's own checks
        r += 1
        for ci in range(1, len(cols) + 1):
            cell = ws.cell(row=r, column=ci, value=None)
            cell.font = font(10)
            cell.alignment = WRAP_TOP
            cell.border = BORDER
    last = r
    ncol = get_column_letter(len(cols))
    list_validation(ws, STATUSES, f"H2:H{last}")
    list_validation(ws, PRIORITIES, f"F2:F{last}")
    list_validation(ws, [e for e, _ in EVIDENCE], f"G2:G{last}")
    status_cf(ws, f"H2:H{last}")
    priority_cf(ws, f"F2:F{last}")
    banding_cf(ws, f"A2:{ncol}{last}")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{ncol}{last}"
    if tab_color:
        ws.sheet_properties.tabColor = tab_color
    print_setup(ws)
    return ws, last_check, last


# --------------------------------------------------------------------------
# Start here tab
# --------------------------------------------------------------------------
def build_start(wb, blocks, tab_color=C["accent"]):
    """blocks: list of tuples. Types: title, subtitle, h, p, steps, kv (list of (k, v)),
    status_legend, priority_legend, evidence_legend, links (list of (label, url)), spacer."""
    ws = wb.active
    ws.title = "Start here"
    ws.sheet_properties.tabColor = tab_color
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 92
    r = 1

    def para(text, size=10, bold=False, color=None, italic=False, height=None):
        nonlocal r
        r += 1
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        c = ws.cell(row=r, column=2, value=text)
        c.font = font(size, bold, color, italic)
        c.alignment = WRAP_TOP
        ws.row_dimensions[r].height = height or (max(1, est_lines(text, 112)) * (size + 5) + 4)

    def kv(k, v, kfill=None, kfont=None, link=None):
        nonlocal r
        r += 1
        a = ws.cell(row=r, column=2, value=k)
        a.font = kfont or font(10, True)
        a.alignment = WRAP_TOP
        if kfill:
            a.fill = fill(kfill)
        b = ws.cell(row=r, column=3, value=link_formula(link, v) if link else v)
        b.font = font(10, color=C["accent"], underline="single") if link else font(10)
        b.alignment = WRAP_TOP
        ws.row_dimensions[r].height = max(est_lines(k, 26), est_lines(v, 90)) * 15 + 4

    for blk in blocks:
        kind = blk[0]
        if kind == "title":
            para(blk[1], 18, True, C["header"], height=30)
        elif kind == "subtitle":
            para(blk[1], 11, False, C["muted"])
        elif kind == "h":
            r += 1
            para(blk[1], 12, True, C["accent"], height=22)
        elif kind == "p":
            para(blk[1])
        elif kind == "note":
            para(blk[1], 9, False, C["muted"], italic=True)
        elif kind == "steps":
            for i, s in enumerate(blk[1], start=1):
                kv(f"Step {i}", s)
        elif kind == "kv":
            for k, v in blk[1]:
                kv(k, v)
        elif kind == "status_legend":
            kv("Not started", "Default for every check.", kfont=Font(name=FONT, color=C["muted"]))
            kv("In progress", "Someone is working on it.", kfill=C["prog_fill"], kfont=Font(name=FONT, bold=True, color=C["prog_font"]))
            kv("Done", "Checked, and fixed if it needed fixing.", kfill=C["done_fill"], kfont=Font(name=FONT, bold=True, color=C["done_font"]))
            kv("N/A", "Does not apply to this site. N/A rows are left out of the % complete.", kfill=C["na_fill"], kfont=Font(name=FONT, color=C["na_font"]))
        elif kind == "priority_legend":
            kv("High", blk[1][0], kfill=C["high"], kfont=Font(name=FONT, bold=True, color=C["header"]))
            kv("Medium", blk[1][1])
            kv("Low", blk[1][2], kfont=Font(name=FONT, color=C["muted"]))
        elif kind == "evidence_legend":
            for k, v in EVIDENCE:
                kv(k, v)
        elif kind == "links":
            for label, url in blk[1]:
                kv(label[0], label[1], link=url)
        elif kind == "spacer":
            r += 1
    print_setup(ws, landscape=False)
    return ws


# --------------------------------------------------------------------------
# Progress tab (COUNTIF / COUNTIFS only, so it works in Sheets and Excel)
# --------------------------------------------------------------------------
def build_progress(wb, sheet, areas, area_label="Area", last_row=500, title="Progress"):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = C["high"]
    s = f"'{sheet}'"
    A, P, S, ID = f"{s}!$B$2:$B${last_row}", f"{s}!$F$2:$F${last_row}", f"{s}!$H$2:$H${last_row}", f"{s}!$C$2:$C${last_row}"
    widths = [3, 30, 10, 10, 13, 13, 10, 12, 28, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws["B2"] = "Progress"
    ws["B2"].font = font(18, True, C["header"])
    ws["B3"] = "Updates automatically as you change the Status column. Formulas use COUNTIF / COUNTIFS, so they work in Google Sheets and Excel."
    ws["B3"].font = font(9, color=C["muted"], italic=True)

    # Overall
    labels = [("Checks", f"=COUNTA({ID})"), ("Done", f'=COUNTIF({S},"Done")'),
              ("In progress", f'=COUNTIF({S},"In progress")'), ("Not started", f'=COUNTIF({S},"Not started")'),
              ("N/A", f'=COUNTIF({S},"N/A")')]
    r = 5
    for k, f_ in labels:
        ws.cell(row=r, column=2, value=k).font = font(10, True)
        c = ws.cell(row=r, column=3, value=f_)
        c.font = font(10)
        c.alignment = Alignment(horizontal="right")
        r += 1
    ws.cell(row=r, column=2, value="% complete (Done / applicable)").font = font(10, True)
    pc = ws.cell(row=r, column=3, value="=IFERROR(C6/(C5-C9),0)")
    pc.number_format = "0%"
    pc.font = font(12, True, C["accent"])
    bar = ws.cell(row=r, column=4, value=f'=REPT("█",ROUND(C{r}*20,0))&REPT("░",20-ROUND(C{r}*20,0))')
    bar.font = Font(name=FONT, size=11, color=C["accent"])
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=8)
    pct_row = r
    r += 1
    ws.cell(row=r, column=2, value="Open High-priority checks").font = font(10, True)
    c = ws.cell(row=r, column=3, value=f'=COUNTIFS({P},"High",{S},"<>Done",{S},"<>N/A")')
    c.font = font(12, True, C["header"])
    c.fill = fill(C["high"])

    # By area
    r += 2
    hdr = [area_label, "Checks", "Done", "In progress", "Not started", "N/A", "% done", "Open High", ]
    for i, h in enumerate(hdr, start=2):
        c = ws.cell(row=r, column=i, value=h)
        c.font = font(10, True, C["white"])
        c.fill = fill(C["header"])
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[r].height = 28
    first = r + 1
    for a in areas:
        r += 1
        ws.cell(row=r, column=2, value=a).font = font(10, True)
        ref = f"$B{r}"
        fs = [f"=COUNTIF({A},{ref})",
              f'=COUNTIFS({A},{ref},{S},"Done")', f'=COUNTIFS({A},{ref},{S},"In progress")',
              f'=COUNTIFS({A},{ref},{S},"Not started")', f'=COUNTIFS({A},{ref},{S},"N/A")',
              f"=IFERROR(D{r}/(C{r}-G{r}),0)",
              f'=COUNTIFS({A},{ref},{P},"High",{S},"<>Done",{S},"<>N/A")']
        for i, f_ in enumerate(fs, start=3):
            c = ws.cell(row=r, column=i, value=f_)
            c.font = font(10)
            c.alignment = Alignment(horizontal="right")
            if i == 8:
                c.number_format = "0%"
        bar = ws.cell(row=r, column=10, value=f'=REPT("█",ROUND(H{r}*10,0))&REPT("░",10-ROUND(H{r}*10,0))')
        bar.font = Font(name=FONT, size=10, color=C["accent"])
    ws.column_dimensions["J"].width = 16
    ws.conditional_formatting.add(f"B{first}:J{r}", FormulaRule(formula=["MOD(ROW(),2)=0"], fill=fill(C["band"])))
    ws.conditional_formatting.add(f"I{first}:I{r}", CellIsRule(operator="greaterThan", formula=["0"],
                                                               fill=fill(C["high"]), font=Font(name=FONT, bold=True, color=C["header"])))

    # By priority
    r += 2
    for i, h in enumerate(["Priority", "Checks", "Done", "Open", "% done"], start=2):
        c = ws.cell(row=r, column=i, value=h)
        c.font = font(10, True, C["white"])
        c.fill = fill(C["header"])
    for p in PRIORITIES:
        r += 1
        ws.cell(row=r, column=2, value=p).font = font(10, True)
        ref = f"$B{r}"
        fs = [f"=COUNTIF({P},{ref})", f'=COUNTIFS({P},{ref},{S},"Done")',
              f'=COUNTIFS({P},{ref},{S},"<>Done",{S},"<>N/A")',
              f'=IFERROR(D{r}/(C{r}-COUNTIFS({P},{ref},{S},"N/A")),0)']
        for i, f_ in enumerate(fs, start=3):
            c = ws.cell(row=r, column=i, value=f_)
            c.font = font(10)
            c.alignment = Alignment(horizontal="right")
            if i == 6:
                c.number_format = "0%"
    priority_cf(ws, f"B{r-2}:B{r}")
    print_setup(ws, landscape=False)
    return ws, r


# --------------------------------------------------------------------------
# Simple log / tracker tabs
# --------------------------------------------------------------------------
def build_log(wb, title, columns, rows=200, validations=None, intro=None, example_rows=None, tab_color=None,
              formulas=None, wrap_cols=None):
    """columns: list of (header, width). validations: {col_letter: [values]}.
    formulas: {col_letter: template with {r}} filled for every row."""
    ws = wb.create_sheet(title)
    start = 1
    if intro:
        ws.cell(row=1, column=1, value=intro).font = font(9, color=C["muted"], italic=True)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
        ws.row_dimensions[1].height = max(1, est_lines(intro, sum(w for _, w in columns))) * 14 + 6
        ws.cell(row=1, column=1).alignment = WRAP_TOP
        start = 2
    header_row(ws, [h for h, _ in columns], row=start, widths=[w for _, w in columns])
    first = start + 1
    last = start + rows
    example_rows = example_rows or []
    for i in range(rows):
        r = first + i
        ex = example_rows[i] if i < len(example_rows) else None
        for ci in range(1, len(columns) + 1):
            col = get_column_letter(ci)
            v = None
            if formulas and col in formulas:
                v = formulas[col].format(r=r)
            elif ex is not None and ci - 1 < len(ex):
                v = ex[ci - 1]
            cell = ws.cell(row=r, column=ci, value=v if v != "" else None)
            cell.font = font(10, italic=ex is not None and not (formulas and col in formulas),
                             color=C["muted"] if ex is not None else None)
            cell.alignment = WRAP_TOP if (wrap_cols is None or col in wrap_cols) else Alignment(vertical="top")
            cell.border = BORDER
    for col, values in (validations or {}).items():
        list_validation(ws, values, f"{col}{first}:{col}{last}")
    ncol = get_column_letter(len(columns))
    ws.freeze_panes = ws.cell(row=first, column=2)
    ws.auto_filter.ref = f"A{start}:{ncol}{last}"
    if tab_color:
        ws.sheet_properties.tabColor = tab_color
    print_setup(ws)
    return ws, first, last


# --------------------------------------------------------------------------
# Markdown version
# --------------------------------------------------------------------------
def md_link(label, url):
    return f"[{label}]({url})" if url else ""


def write_markdown(path, title, intro_md, checks, area_order, outro_md, area_intros=None):
    out = [f"# {title}", "", intro_md.strip(), ""]
    for area in area_order:
        out.append(f"## {area}")
        out.append("")
        if area_intros and area in area_intros:
            out += [area_intros[area], ""]
        for ck in [c for c in checks if c["area"] == area]:
            out.append(f"- [ ] **{ck['id']} · {ck['check']}**  ")
            out.append(f"  *Priority: {ck['priority']} · Evidence: {ck['evidence']}*  ")
            out.append(f"  **Why it matters:** {ck['why']}  ")
            out.append(f"  **How to check:** {ck['how']}  ")
            links = [x for x in (md_link(ck.get('learn_label'), ck.get('learn_url')),
                                 md_link(ck.get('source_label'), ck.get('source_url'))) if x]
            if links:
                out.append(f"  **Learn more:** {' · '.join(links)}")
            out.append("")
    out += [outro_md.strip(), ""]
    pathlib.Path(path).write_text("\n".join(out), encoding="utf-8")


# ===========================================================================
# GEO checklist: folder-specific content and tabs
# Run:  python build.py   (needs openpyxl)  ->  geo-checklist.xlsx + checklist.md
# ===========================================================================
TITLE = "Generative Engine Optimisation (GEO) Checklist"
XLSX = "geo-checklist.xlsx"
AREAS = ['1. Crawl access', '2. Indexing & eligibility', '3. Entity clarity', '4. Chunk extractability', '5. Claim quality', '6. Structured data', '7. Trust & corroboration', '8. Measurement']
BOTS = [
    ('Googlebot', 'Google', 'Google Search, including AI Overviews and AI Mode', 'Blocking removes pages from Google Search and its AI features.', 'Respects robots.txt', 'https://developers.google.com/search/docs/appearance/ai-features', 'Google: AI features and your website'),
    ('Google-Extended', 'Google', 'Control token only (no separate crawler) for Gemini Apps and Vertex AI training and grounding', 'Does not affect Google Search or AI Overviews.', 'robots.txt token', 'https://developers.google.com/search/docs/crawling-indexing/google-common-crawlers', 'Google: common crawlers (Google-Extended)'),
    ('Bingbot', 'Microsoft', "Bing search; Bing's index also grounds Microsoft Copilot answers", 'Blocking removes pages from Bing and the answers built on it.', 'Respects robots.txt', 'https://www.bing.com/webmasters/help/which-crawlers-does-bing-use-8c184ec0', 'Bing: which crawlers Bing uses'),
    ('OAI-SearchBot', 'OpenAI', 'Shows and links sites in ChatGPT search', "Sites opted out won't be shown in ChatGPT search answers (may still appear as navigational links).", 'Respects robots.txt; changes take about 24 hours', 'https://platform.openai.com/docs/bots', 'OpenAI: crawlers and user agents'),
    ('GPTBot', 'OpenAI', "Crawls content that may be used to train OpenAI's foundation models", 'Signals content should not be used for training. Independent of OAI-SearchBot.', 'Respects robots.txt', 'https://platform.openai.com/docs/bots', 'OpenAI: crawlers and user agents'),
    ('ChatGPT-User', 'OpenAI', 'Visits pages when a ChatGPT user or Custom GPT asks it to', 'User-initiated; OpenAI says robots.txt rules may not apply. Not used for search inclusion.', 'May not apply', 'https://platform.openai.com/docs/bots', 'OpenAI: crawlers and user agents'),
    ('Claude-SearchBot', 'Anthropic', "Indexes content to improve Claude's search results", "Blocking may reduce visibility in Claude's search answers.", 'Respects robots.txt', 'https://support.claude.com/en/articles/8896518-does-anthropic-crawl-data-from-the-web-and-how-can-site-owners-block-the-crawler', 'Anthropic: crawlers and opt-out'),
    ('ClaudeBot', 'Anthropic', 'Collects content that could contribute to model training', 'Signals future content should be excluded from training.', 'Respects robots.txt', 'https://support.claude.com/en/articles/8896518-does-anthropic-crawl-data-from-the-web-and-how-can-site-owners-block-the-crawler', 'Anthropic: crawlers and opt-out'),
    ('Claude-User', 'Anthropic', 'Fetches pages when a Claude user asks a question', 'Blocking stops Claude retrieving your pages for users.', 'Respects robots.txt', 'https://support.claude.com/en/articles/8896518-does-anthropic-crawl-data-from-the-web-and-how-can-site-owners-block-the-crawler', 'Anthropic: crawlers and opt-out'),
    ('PerplexityBot', 'Perplexity', 'Surfaces and links sites in Perplexity search; not used for foundation-model training', "Blocking removes you from Perplexity's search results.", 'Respects robots.txt; up to 24 hours to update', 'https://docs.perplexity.ai/guides/bots', 'Perplexity: crawlers'),
    ('Perplexity-User', 'Perplexity', "Fetches pages for a user's question", 'User-initiated; Perplexity says it generally ignores robots.txt.', 'Generally ignores robots.txt', 'https://docs.perplexity.ai/guides/bots', 'Perplexity: crawlers'),
    ('Applebot', 'Apple', 'Siri, Spotlight and other Apple search features', "Blocking removes pages from Apple's search features.", 'Respects robots.txt', 'https://support.apple.com/en-us/119829', 'Apple: about Applebot'),
    ('Applebot-Extended', 'Apple', "Control token for using Applebot-crawled content to train Apple's generative models", 'Opts out of training only.', 'robots.txt token', 'https://support.apple.com/en-us/119829', 'Apple: about Applebot'),
    ('Meta-ExternalAgent', 'Meta', 'Crawls for uses such as training AI models or indexing content', 'Opts out of that crawling.', "Check Meta's docs", 'https://developers.facebook.com/docs/sharing/webmasters/web-crawlers', 'Meta: web crawlers'),
    ('CCBot', 'Common Crawl', 'Builds the open Common Crawl dataset, widely used in model training', 'Keeps new content out of future Common Crawl snapshots.', 'Respects robots.txt', 'https://commoncrawl.org/ccbot', 'Common Crawl: CCBot'),
]
ASSISTANTS = ["ChatGPT", "Perplexity", "Claude", "Gemini", "Google AI Mode", "Google AI Overview", "Microsoft Copilot", "Other"]

START = [
    ("title", TITLE),
    ("subtitle", "44 checks to help AI search systems find, understand, trust and cite your content. Free from MLforSEO, October 2026."),
    ("h", "What this is"),
    ("p", "A working checklist for getting content into AI answers (Google AI Overviews and AI Mode, ChatGPT, Perplexity, Claude, Gemini, Copilot). "
          "It runs in the order AI systems process content: can they reach it, can they find and recognise it, can they lift a clean passage out of it, "
          "do they believe it, and can you see the result. Every check links to the MLforSEO article that explains it, plus the official documentation where one exists."),
    ("p", "Generative engine optimisation (GEO) is mostly good SEO done with a different reader in mind. Google says there are no extra requirements to appear in its AI features beyond "
          "being indexed and eligible for a snippet, so the foundations come first. Each check carries an Evidence label so you can see which items platforms document, "
          "which come from research, and which are MLforSEO practice."),
    ("h", "How to use it"),
    ("steps", [
        "Make your own copy (File > Make a copy in Google Sheets). Don't request edit access to the original.",
        "Work through the Checklist tab area by area. Set Status for each check, add an Owner and write what you found in Notes.",
        "Mark checks that don't apply as N/A. They are left out of the % complete.",
        "Use the filter arrows in the header row, for example Priority = High and Status = Not started, to build your first sprint.",
        "Fill in the AI crawlers tab while you work on area 1, and start the Prompt tracker for area 8. The Progress tab calculates everything.",
        "Rerun the Prompt tracker monthly. AI answers change between runs, so look at the trend, not one result.",
    ]),
    ("h", "Tabs"),
    ("kv", [
        ("Checklist", "The 44 checks, with why each matters, how to check it, priority, evidence level, status, owner, notes and links. Blank rows at the bottom are for your own checks."),
        ("AI crawlers", "Every major AI crawler and user-agent token, what it does, what blocking it does, and a column for your decision. Checked against each provider's documentation in October 2026."),
        ("Prompt tracker", "Log the customer questions you test in each assistant and whether you were mentioned, cited and described accurately."),
        ("Progress", "Completion by area and priority, open High-priority checks, and mention, citation and accuracy rates from the Prompt tracker."),
    ]),
    ("h", "Status"),
    ("status_legend",),
    ("h", "Priority"),
    ("priority_legend", ["Fix first. Blocks access, eligibility or entity recognition, or misrepresents you.",
                         "Do next. Improves how often and how well you are selected.",
                         "Nice to have, or an early-warning signal to monitor."]),
    ("h", "Evidence levels"),
    ("evidence_legend",),
    ("h", "What this checklist can't promise"),
    ("p", "No checklist guarantees citations. AI answers vary by prompt wording, location, account and run, and providers change their systems often. "
          "Use the checks to remove the reasons you would be skipped, then measure what happens. Facts were checked in October 2026; recheck the official sources linked in each row before big decisions."),
    ("h", "Credits"),
    ("links", [
        (("Created by", "Lazarina Stoy · MLforSEO"), "https://www.mlforseo.com/experts/lazarina-stoy/"),
        (("Built on", "Beatrice Gamba's AI search series on MLforSEO"), "https://www.mlforseo.com/experts/beatrice-gamba/"),
        (("Go deeper", "AI Search Optimisation & Agentic SEO course"), "https://www.mlforseo.com/courses/ai-search-optimisation-agentic-seo/"),
        (("More free resources", "MLforSEO resources library"), "https://www.mlforseo.com/resources-library/"),
        (("Source files", "GitHub: lazarinastoy/mlforseo-code"), "https://github.com/lazarinastoy/mlforseo-code/tree/main/geo-checklist"),
    ]),
    ("note", "Version 1.0, October 2026. Free to use and adapt (MIT licence); please keep the credit."),
]


def build_bots(wb):
    cols = [("User-agent token", 20), ("Company", 13), ("What it does", 40), ("What blocking it does", 40),
            ("robots.txt", 20), ("Docs", 28), ("Your decision", 14), ("In your robots.txt now", 18), ("Notes", 30)]
    ws, first, last = build_log(wb, "AI crawlers", cols, rows=len(BOTS) + 5,
                                intro="Tokens and behaviour as documented by each provider, checked October 2026. Search and user-fetch bots affect whether you can be shown in AI answers; training bots don't. Record your policy in 'Your decision' and what robots.txt actually says in the next column.",
                                validations={"G": ["Allow", "Disallow", "Partial"], "H": ["Allowed", "Blocked", "Partly blocked", "Not mentioned"]},
                                tab_color=C["accent"])
    for i, (tok, co, what, block, rob, url, label) in enumerate(BOTS):
        r = first + i
        vals = [tok, co, what, block, rob, link_formula(url, label)]
        for ci, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=ci, value=v)
            c.font = font(10, bold=ci == 1, color=C["accent"] if ci == 6 else None, underline="single" if ci == 6 else None)
        ws.row_dimensions[r].height = max(est_lines(what, 40), est_lines(block, 40)) * 15 + 6
    ws.conditional_formatting.add(f"H{first}:H{last}", CellIsRule(operator="equal", formula=['"Blocked"'], fill=fill(C["warn_fill"]), font=Font(name=FONT, color=C["warn_font"], bold=True)))
    ws.conditional_formatting.add(f"G{first}:G{last}", CellIsRule(operator="equal", formula=['"Allow"'], font=Font(name=FONT, color=C["done_font"], bold=True)))
    banding_cf(ws, f"A{first}:I{last}")
    return ws


def build_prompts(wb, rows=300):
    cols = [("Date", 11), ("Prompt (keep wording fixed)", 40), ("Topic / intent", 18), ("Assistant", 16),
            ("Mentioned?", 11), ("Cited with link?", 11), ("Cited URL", 30), ("Described accurately?", 13),
            ("How you appear", 20), ("Competitors named", 26), ("Notes", 30)]
    examples = [["2026-10-01", "What is the best tool for [your category] for small teams?", "Commercial", "ChatGPT", "Yes", "No", "", "Partly", "Listed among several", "Competitor A, Competitor B", "EXAMPLE ROW: delete. Price quoted was last year's."],
                ["2026-10-01", "What is the best tool for [your category] for small teams?", "Commercial", "Perplexity", "Yes", "Yes", "https://www.example.com/pricing/", "Yes", "Primary recommendation", "Competitor A", "EXAMPLE ROW: delete."]]
    ws, first, last = build_log(wb, "Prompt tracker", cols, rows=rows,
                                intro="One row per prompt, per assistant, per run. Use 10 to 50 real customer questions and rerun them monthly with the same wording, logged out where possible. Delete the two example rows before you start.",
                                validations={"D": ASSISTANTS, "E": ["Yes", "No"], "F": ["Yes", "No"], "H": ["Yes", "Partly", "No", "Not mentioned"],
                                             "I": ["Primary recommendation", "Listed among several", "Passing mention", "Not present"]},
                                example_rows=examples, tab_color=C["accent"])
    for col, good, bad in (("E", "Yes", "No"), ("F", "Yes", "No")):
        ws.conditional_formatting.add(f"{col}{first}:{col}{last}", CellIsRule(operator="equal", formula=[f'"{good}"'], fill=fill(C["done_fill"]), font=Font(name=FONT, color=C["done_font"], bold=True)))
    ws.conditional_formatting.add(f"H{first}:H{last}", CellIsRule(operator="equal", formula=['"No"'], fill=fill(C["warn_fill"]), font=Font(name=FONT, color=C["warn_font"], bold=True)))
    ws.conditional_formatting.add(f"H{first}:H{last}", CellIsRule(operator="equal", formula=['"Partly"'], fill=fill(C["high"])))
    banding_cf(ws, f"A{first}:K{last}")
    return ws, first, last


def add_ai_visibility(ws, start_row, first, last):
    s = "'Prompt tracker'"
    P, AS, M, CI, AC = (f"{s}!$B${first}:$B${last}", f"{s}!$D${first}:$D${last}", f"{s}!$E${first}:$E${last}",
                        f"{s}!$F${first}:$F${last}", f"{s}!$H${first}:$H${last}")
    r = start_row + 2
    ws.cell(row=r, column=2, value="AI visibility (from the Prompt tracker)").font = font(12, True, C["accent"])
    r += 1
    hdr = ["Assistant", "Runs logged", "Mentioned", "Mention rate", "Cited with link", "Citation rate", "Accurate", "Accuracy rate"]
    for i, h in enumerate(hdr, start=2):
        c = ws.cell(row=r, column=i, value=h)
        c.font = font(10, True, C["white"])
        c.fill = fill(C["header"])
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[r].height = 28
    first_row = r + 1
    for a in ["All assistants"] + ASSISTANTS:
        r += 1
        ws.cell(row=r, column=2, value=a).font = font(10, True)
        if a == "All assistants":
            fs = [f'=COUNTIF({M},"Yes")+COUNTIF({M},"No")', f'=COUNTIF({M},"Yes")', f"=IFERROR(D{r}/C{r},0)",
                  f'=COUNTIF({CI},"Yes")', f"=IFERROR(F{r}/C{r},0)", f'=COUNTIF({AC},"Yes")', f"=IFERROR(H{r}/D{r},0)"]
        else:
            ref = f"$B{r}"
            fs = [f'=COUNTIFS({AS},{ref},{M},"Yes")+COUNTIFS({AS},{ref},{M},"No")', f'=COUNTIFS({AS},{ref},{M},"Yes")',
                  f"=IFERROR(D{r}/C{r},0)", f'=COUNTIFS({AS},{ref},{CI},"Yes")', f"=IFERROR(F{r}/C{r},0)",
                  f'=COUNTIFS({AS},{ref},{AC},"Yes")', f"=IFERROR(H{r}/D{r},0)"]
        for i, f_ in enumerate(fs, start=3):
            c = ws.cell(row=r, column=i, value=f_)
            c.font = font(10, bold=a == "All assistants")
            c.alignment = Alignment(horizontal="right")
            if i in (5, 7, 9):
                c.number_format = "0%"
    ws.conditional_formatting.add(f"B{first_row}:I{r}", FormulaRule(formula=["MOD(ROW(),2)=0"], fill=fill(C["band"])))
    r += 1
    ws.cell(row=r, column=2, value="Mention rate = runs where you were named / runs logged. Citation rate = runs with a link to your site / runs logged. Accuracy rate = accurate descriptions / runs where you were named.").font = font(9, color=C["muted"], italic=True)
    return r


def main():
    checks = read_checks()
    wb = Workbook()
    build_start(wb, START)
    ws, last_check, last = build_checklist(wb, checks, area_label="Area")
    build_bots(wb)
    pws, pfirst, plast = build_prompts(wb)
    prog, r = build_progress(wb, "Checklist", AREAS, last_row=last)
    add_ai_visibility(prog, r, pfirst, plast)
    wb.save(HERE / XLSX)
    write_markdown(HERE / "checklist.md", TITLE, MD_INTRO, checks, AREAS, MD_OUTRO)
    print("wrote", XLSX, "and checklist.md with", len(checks), "checks")


MD_INTRO = """
*44 checks to help AI search systems find, understand, trust and cite your content. By [Lazarina Stoy](https://www.mlforseo.com/experts/lazarina-stoy/) for [MLforSEO](https://www.mlforseo.com/), built on [Beatrice Gamba](https://www.mlforseo.com/experts/beatrice-gamba/)'s AI search series. Checked October 2026.*

This is the plain-text version of `geo-checklist.xlsx` (the Google Sheets version has status dropdowns, an AI crawler reference tab, a prompt tracker and a progress dashboard). The areas follow the order AI systems process content: reach it, recognise it, extract it, trust it, and then measure what happened.

Google says there are no extra requirements to appear in AI Overviews or AI Mode beyond being indexed and eligible for a snippet, so the foundations come first. Each check has an **evidence level**:

| Evidence | Meaning |
|---|---|
| Google-documented | Stated in Google Search Central or Google Help documentation |
| Platform-documented | Stated in another platform's own docs (OpenAI, Anthropic, Perplexity, Microsoft Bing, Apple, Meta, Wikidata, schema.org) |
| Research-backed | Supported by published research or a large public dataset; may not carry over to every site |
| MLforSEO practice | What we teach and use in audits, based on how retrieval and entity systems work; not confirmed by the platforms |
| Unconfirmed (optional) | A cheap extra no major platform has confirmed using |

**Priority:** High = fix first (blocks access, eligibility or recognition, or misrepresents you). Medium = do next. Low = nice to have or an early-warning signal.
"""

MD_OUTRO = """
---

## What this checklist can't promise

No checklist guarantees citations. AI answers vary by prompt wording, location, account and run, and providers change their systems often. Use the checks to remove the reasons you would be skipped, then measure what happens.

## Go deeper

- [AI Search Optimisation & Agentic SEO course](https://www.mlforseo.com/courses/ai-search-optimisation-agentic-seo/) (Beatrice Gamba, MLforSEO Academy)
- [MLforSEO resources library](https://www.mlforseo.com/resources-library/)

Created by Lazarina Stoy · MLforSEO. Free to use and adapt under the repository's MIT licence; please keep the credit.
"""

if __name__ == "__main__":
    main()
