#!/usr/bin/env python3
"""Build keyword-mapping-template.xlsx (MLforSEO, Lazarina Stoy).

    pip install openpyxl
    python build_template.py

The workbook is plain formulas (INDEX/MATCH, COUNTIFS, SUMIFS, SUMPRODUCT, IF, IFERROR), so it works the same
in Google Sheets and Excel. No XLOOKUP, LET, LAMBDA, FILTER or UNIQUE.
"""
import os
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment

from brand_style import (DARK, F_ACCENT, F_DARK, F_HI, F_LAV, F_RED, band, credits, esc, flag_rule, font,
                         group_label, header_row, legend, lines, list_validation, note, number_validation,
                         sample_banner, title, widths)

OUT = Path(os.environ.get("OUT", Path(__file__).with_name("keyword-mapping-template.xlsx")))
FIRST, LAST = 3, int(os.environ.get("KM_LAST", 1002))          # keyword rows on "Keyword Map"
URL_LAST = int(os.environ.get("URL_LAST", 502))
CL_LAST = int(os.environ.get("CL_LAST", 302))
KM = "'Keyword Map'"
INTENTS = ["Informational", "Commercial", "Transactional", "Navigational"]
POST_URL = "https://www.mlforseo.com/machine-learning-implementation-guides/keyword-research/how-to-turn-semantic-keyword-research-into-an-actionable-seo-workflow/"
COURSE_URL = "https://www.mlforseo.com/courses/semantic-ml-enabled-keyword-research/"

# ---------------------------------------------------------------- settings data
PAGE_TYPES = [  # page type: fit for Informational, Commercial, Transactional, Navigational
    ("Homepage", "Partial", "Partial", "Partial", "Yes"),
    ("Category / collection", "No", "Yes", "Yes", "Partial"),
    ("Product / service page", "No", "Yes", "Yes", "Partial"),
    ("Comparison / review", "Partial", "Yes", "Partial", "No"),
    ("Blog post / guide", "Yes", "Partial", "No", "No"),
    ("Glossary / definition", "Yes", "No", "No", "No"),
    ("Landing page", "No", "Yes", "Yes", "Yes"),
    ("Store / location page", "No", "Partial", "Yes", "Yes"),
    ("Support / docs", "Yes", "No", "No", "Yes"),
    ("Tool / calculator", "Yes", "Partial", "Partial", "No"),
    ("About / company", "Partial", "No", "No", "Yes"),
    ("", "", "", "", ""),
]
# Same word lists as the MLforSEO search-intent-regex-classifier (plurals spelled out, as there is no regex here),
# plus "shop", "sale", "coupon code" for ecommerce. Navigational = your own brand words: edit them.
INTENT_RULES = {
    "Navigational": ["stride", "stride and co", "login", "log in", "sign in", "my account"],
    "Transactional": ["buy", "price", "prices", "pricing", "expensive", "recommendation", "recommendations",
                      "recommended", "near me", "firm", "coupon", "coupons", "coupon code", "order", "purchase",
                      "discount", "deal", "deals", "free trial", "demo", "shop", "sale"],
    "Commercial": ["best", "top", "vs", "versus", "review", "reviews", "cheap", "comparison", "compare",
                   "alternative", "alternatives"],
    "Informational": ["how", "why", "what", "does", "which", "when", "who", "guide", "tutorial", "learn",
                      "example", "examples", "resource", "resources", "ideas", "tips", "meaning", "definition"],
}
RULE_FIRST, RULE_LAST = 5, 124

# ---------------------------------------------------------------- sample data (synthetic)
B = "https://www.stride-example.com"
SAMPLE = [
    # keyword, cluster, volume, kd, ranking url, position, target url, status, page type, role, pasted intent, notes
    ("trail running shoes", "Trail shoes", 9900, 54, "/trail-running-shoes/", 14, "/trail-running-shoes/", "Existing", "Category / collection", "Primary", "Commercial", "Intent pasted from a SERP-based classifier"),
    ("best trail running shoes", "Trail shoes", 4400, 61, "/blog/best-trail-running-shoes/", 9, "/blog/best-trail-running-shoes/", "Existing", "Comparison / review", "Primary", "", ""),
    ("trail running shoes for beginners", "Trail shoes", 590, 22, "/blog/best-trail-running-shoes/", 18, "/blog/trail-running-for-beginners/", "Planned", "Blog post / guide", "Primary", "", "New guide planned"),
    ("waterproof trail running shoes", "Trail shoes", 1300, 35, "/trail-running-shoes/", 21, "/trail-running-shoes/waterproof/", "Planned", "Category / collection", "Primary", "", ""),
    ("buy trail running shoes", "Trail shoes", 210, 40, "", "", "/blog/trail-running-for-beginners/", "Planned", "Blog post / guide", "Secondary", "", ""),
    ("wide running shoes", "Wide fit", 2900, 30, "/wide-fit-running-shoes/", 7, "/wide-fit-running-shoes/", "Existing", "Category / collection", "Primary", "Commercial", ""),
    ("running shoes for wide feet", "Wide fit", 1900, 28, "/blog/wide-feet-guide/", 11, "/wide-fit-running-shoes/", "Existing", "Category / collection", "Secondary", "", ""),
    ("running shoes for wide feet", "Wide fit", 1900, 28, "/blog/wide-feet-guide/", 11, "/blog/wide-feet-guide/", "Existing", "Blog post / guide", "Primary", "", "Also mapped to the category page"),
    ("how to choose running shoes for wide feet", "Wide fit", 320, 12, "/blog/wide-feet-guide/", 4, "/blog/wide-feet-guide/", "Existing", "Blog post / guide", "Secondary", "", ""),
    ("marathon running shoes", "Marathon", 2400, 48, "/marathon-running-shoes/", 16, "/marathon-running-shoes/", "Existing", "Category / collection", "Primary", "Commercial", ""),
    ("best marathon shoes", "Marathon", 1600, 52, "/marathon-running-shoes/", 19, "/marathon-running-shoes/", "Existing", "Category / collection", "Primary", "", ""),
    ("carbon plate running shoes", "Marathon", 1000, 44, "", "", "/marathon-running-shoes/carbon-plate/", "Planned", "Category / collection", "Primary", "", ""),
    ("what are carbon plate shoes", "Marathon", 480, 18, "", "", "/glossary/carbon-plate/", "Planned", "Glossary / definition", "Primary", "", ""),
    ("hoka vs brooks", "Brand comparisons", 880, 26, "/blog/hoka-vs-brooks/", 6, "/blog/hoka-vs-brooks/", "Existing", "Comparison / review", "Primary", "", ""),
    ("hoka clifton vs brooks ghost", "Brand comparisons", 720, 21, "/blog/hoka-vs-brooks/", 8, "/blog/hoka-vs-brooks/", "Existing", "Comparison / review", "Secondary", "", ""),
    ("asics gel nimbus review", "Brand comparisons", 1300, 33, "", "", "/reviews/asics-gel-nimbus/", "Planned", "Comparison / review", "Primary", "", ""),
    ("stride and co discount code", "Brand & store", 390, 5, "/offers/", 2, "/offers/", "Existing", "Landing page", "Primary", "", ""),
    ("stride and co returns", "Brand & store", 170, 3, "/help/returns/", 1, "/help/returns/", "Existing", "Support / docs", "Primary", "", ""),
    ("running shoe shop london", "Brand & store", 880, 38, "/stores/london/", 5, "/stores/london/", "Existing", "Store / location page", "Primary", "", ""),
    ("gait analysis near me", "Gait analysis", 1000, 30, "/stores/london/", 7, "/stores/london/", "Existing", "Store / location page", "Secondary", "", ""),
    ("gait analysis", "Gait analysis", 2400, 36, "/services/gait-analysis/", 9, "/services/gait-analysis/", "Existing", "Product / service page", "Primary", "", ""),
    ("running shoes sale", "Deals", 1600, 41, "/sale/", 3, "/sale/", "Existing", "Category / collection", "Primary", "", ""),
    ("cheap running shoes", "Deals", 2900, 45, "/sale/", 12, "/sale/", "Existing", "Category / collection", "Secondary", "", ""),
    ("how often should you replace running shoes", "Shoe care", 1900, 15, "/blog/replace-running-shoes/", 3, "/blog/replace-running-shoes/", "Existing", "Blog post / guide", "Primary", "", ""),
    ("how to clean running shoes", "Shoe care", 1600, 19, "", "", "/blog/clean-running-shoes/", "Planned", "Blog post / guide", "Primary", "", ""),
    ("running shoe lifespan", "Shoe care", 260, 10, "/blog/replace-running-shoes/", 5, "/blog/replace-running-shoes/", "Existing", "Blog post / guide", "Secondary", "informational", "Intent pasted from the regex classifier"),
    ("zero drop running shoes", "Shoe types", 880, 27, "", "", "", "Not mapped yet", "", "", "", "Decide: new category or blog post?"),
]


def full(u):
    return B + u if u else ""


def build():
    wb = Workbook()
    start = wb.active
    start.title = "Start here"
    km = wb.create_sheet("Keyword Map")
    us = wb.create_sheet("URL Summary")
    cs = wb.create_sheet("Cluster Summary")
    dash = wb.create_sheet("Dashboard")
    st = wb.create_sheet("Settings")

    # ============================================================ Settings
    title(st, "A1", "Settings: edit the rules, the sheet recalculates")
    note(st, "A2", "Yellow cells are yours to change. Page types feed the dropdown on Keyword Map.")
    st["A3"] = "Thresholds and weights"; st["A3"].font = font(11, True, DARK)
    settings = [
        ("Max primary keywords per URL", 1, "More than this on one URL raises the 'Too many primary keywords' flag. One primary keyword (or cluster head) per page is my default."),
        ("Volume points (max)", 50, "Points a keyword gets for search volume (log scale)."),
        ("Difficulty points (max)", 25, "Points for ease: KD 0 gets all of them, KD 100 gets none. A blank KD counts as 50."),
        ("Volume that earns full volume points", 5000, "Keywords at or above this volume get all the volume points."),
    ]
    for i, (k, v, why) in enumerate(settings, start=4):
        st.cell(row=i, column=1, value=k).font = font(10)
        c = st.cell(row=i, column=2, value=v); c.fill = F_HI; c.font = font(10, True)
        st.cell(row=i, column=3, value=why).font = font(9, italic=True, color="6B6880")
    st["A9"] = "Intent"; st["B9"] = "Priority points"
    for c in ("A9", "B9"):
        st[c].fill = F_DARK; st[c].font = font(10, True, "FFFFFF")
    for i, (k, v) in enumerate([("Transactional", 25), ("Commercial", 20), ("Informational", 10), ("Navigational", 5), ("Unclassified", 0)], start=10):
        st.cell(row=i, column=1, value=k).font = font(10)
        c = st.cell(row=i, column=2, value=v); c.fill = F_HI
    # page type x intent fit matrix (D4:H16)
    st["D3"] = "Page type × intent fit (Yes / Partial / No). 'No' raises the intent-mismatch flag."; st["D3"].font = font(11, True, DARK)
    for j, h in enumerate(["Page type"] + INTENTS):
        c = st.cell(row=4, column=4 + j, value=h); c.fill = F_DARK; c.font = font(10, True, "FFFFFF")
    for i, row in enumerate(PAGE_TYPES, start=5):
        for j, v in enumerate(row):
            c = st.cell(row=i, column=4 + j, value=v or None)
            c.fill = F_HI if j == 0 else F_LAV
            c.font = font(10, j == 0)
    list_validation(st, "E5:H16", '"Yes,Partial,No"')
    # intent rules (K4:L124)
    st["K3"] = "Intent rules: whole words or phrases, checked in this order: Navigational, Transactional, Commercial, Informational."; st["K3"].font = font(11, True, DARK)
    for j, h in enumerate(["Intent", "Word or phrase"]):
        c = st.cell(row=4, column=11 + j, value=h); c.fill = F_DARK; c.font = font(10, True, "FFFFFF")
    r = RULE_FIRST
    for intent, words in INTENT_RULES.items():
        for w in words:
            st.cell(row=r, column=11, value=intent); st.cell(row=r, column=12, value=w)
            r += 1
    list_validation(st, f"K{RULE_FIRST}:K{RULE_LAST}", '"' + ",".join(INTENTS) + '"')
    band(st, f"K{RULE_FIRST}:L{RULE_LAST}", "K", RULE_FIRST)
    note(st, "K126", "Add rows anywhere up to row 124. Matching is case-insensitive and whole-word ('top' won't match 'laptop').")
    widths(st, {"A": 36, "B": 14, "C": 70, "D": 26, "E": 14, "F": 14, "G": 14, "H": 14, "I": 3, "J": 3, "K": 16, "L": 24})
    st.freeze_panes = "A5"

    # ============================================================ Keyword Map
    group_label(km, 1, 1, 12, "PASTE OR TYPE HERE (keyword research export + your mapping decisions)", "input")
    group_label(km, 1, 13, 28, "CALCULATED: flags update automatically", "calc")
    inputs = ["Keyword", "Cluster", "Search volume", "Keyword difficulty", "Current ranking URL", "Current position",
              "Target URL", "Page status", "Page type", "Keyword role", "Intent (pasted, optional)", "Notes"]
    calcs = ["Keyword (clean)", "Rule-based intent", "Final intent", "Priority score", "Other URLs for this keyword",
             "Cannibalisation", "Primary keywords on target URL", "Too many primaries on URL", "URL spans 2+ clusters",
             "Intent fit (page type)", "Intent mismatch", "Ranking URL is not the target", "Flag count", "Flags",
             "URL index (helper)", "Cluster index (helper)"]
    header_row(km, 2, inputs, 1, "input")
    header_row(km, 2, calcs, 13, "calc")

    rule_k = f"Settings!$K${RULE_FIRST}:$K${RULE_LAST}"
    rule_l = f"Settings!$L${RULE_FIRST}:$L${RULE_LAST}"

    def intent_hits(row, intent):
        return (f'SUMPRODUCT(ISNUMBER(SEARCH(" "&{rule_l}&" "," "&M{row}&" "))*({rule_k}="{intent}")*({rule_l}<>""))')

    rng = lambda col: f"${col}${FIRST}:${col}${LAST}"
    for r in range(FIRST, LAST + 1):
        def norm(ref):
            t = f"TRIM({ref})"
            return f'LOWER(IF(RIGHT({t},1)="/",LEFT({t},LEN({t})-1),{t}))'
        f = {
            "M": f'=IF(A{r}="","",LOWER(TRIM(SUBSTITUTE(SUBSTITUTE(A{r},"?",""),"*",""))))',
            "N": (f'=IF(M{r}="","",IF({intent_hits(r, "Navigational")}>0,"Navigational",IF({intent_hits(r, "Transactional")}>0,"Transactional",'
                  f'IF({intent_hits(r, "Commercial")}>0,"Commercial",IF({intent_hits(r, "Informational")}>0,"Informational","Unclassified")))))'),
            "O": f'=IF(A{r}="","",IF(K{r}<>"",PROPER(K{r}),N{r}))',
            "P": (f'=IF(A{r}="","",ROUND(Settings!$B$5*MIN(1,LOG10(N(C{r})+1)/LOG10(Settings!$B$7+1))'
                  f'+Settings!$B$6*(1-MIN(100,IF(D{r}="",50,N(D{r})))/100)'
                  f'+IFERROR(INDEX(Settings!$B$10:$B$14,MATCH(O{r},Settings!$A$10:$A$14,0)),0),0))'),
            "Q": f'=IF(OR(A{r}="",G{r}=""),"",COUNTIFS({rng("M")},M{r},{rng("G")},"<>"&{esc(f"G{r}")},{rng("G")},"?*"))',
            "R": f'=IF(Q{r}="","",IF(Q{r}>0,"Yes","No"))',
            "S": f'=IF(G{r}="","",COUNTIFS({rng("G")},{esc(f"G{r}")},{rng("J")},"Primary"))',
            "T": f'=IF(S{r}="","",IF(S{r}>Settings!$B$4,"Yes","No"))',
            "U": f'=IF(OR(G{r}="",B{r}=""),"",IF(COUNTIFS({rng("G")},{esc(f"G{r}")},{rng("B")},"<>"&{esc(f"B{r}")},{rng("B")},"?*")>0,"Yes","No"))',
            "V": (f'=IF(OR(G{r}="",I{r}="",O{r}="",O{r}="Unclassified"),"",IFERROR(INDEX(Settings!$E$5:$H$16,'
                  f'MATCH(I{r},Settings!$D$5:$D$16,0),MATCH(O{r},Settings!$E$4:$H$4,0)),"Check"))'),
            "W": f'=IF(V{r}="","",IF(V{r}="No","Yes","No"))',
            "X": f'=IF(OR(E{r}="",G{r}=""),"",IF({norm(f"E{r}")}<>{norm(f"G{r}")},"Yes","No"))',
            "Y": f'=IF(A{r}="","",(N(Q{r})>0)+(T{r}="Yes")+(U{r}="Yes")+(W{r}="Yes")+(X{r}="Yes"))',
            "Z": (f'=IF(A{r}="","",IF(G{r}="","Not mapped",IF(Y{r}=0,"OK",MID(IF(N(Q{r})>0,"; Cannibalisation","")'
                  f'&IF(T{r}="Yes","; Too many primary keywords on URL","")&IF(U{r}="Yes","; URL spans 2+ clusters","")'
                  f'&IF(W{r}="Yes","; Intent mismatch","")&IF(X{r}="Yes","; Ranking URL is not the target",""),3,300))))'),
            "AA": f'=IF(G{r}="","",IF(COUNTIF($G${FIRST}:G{r},{esc(f"G{r}")})=1,MAX($AA${FIRST - 1}:AA{r - 1})+1,""))',
            "AB": f'=IF(B{r}="","",IF(COUNTIF($B${FIRST}:B{r},{esc(f"B{r}")})=1,MAX($AB${FIRST - 1}:AB{r - 1})+1,""))',
        }
        for col, formula in f.items():
            km[f"{col}{r}"] = formula
    # helper header cells must not be numbers (MAX over them must ignore them): they are text, fine.
    for i, row in enumerate(SAMPLE):
        r = FIRST + i
        vals = list(row)
        vals[4] = full(vals[4]); vals[6] = full(vals[6])
        for j, v in enumerate(vals):
            if v != "":
                km.cell(row=r, column=1 + j, value=v)
    sample_banner(km, "M1", "SAMPLE DATA rows 3-29 (synthetic site stride-example.com). Clear columns A-L and paste your own.")
    list_validation(km, f"H{FIRST}:H{LAST}", '"Existing,Planned,Not mapped yet"')
    list_validation(km, f"I{FIRST}:I{LAST}", "=Settings!$D$5:$D$16")
    list_validation(km, f"J{FIRST}:J{LAST}", '"Primary,Secondary"')
    list_validation(km, f"K{FIRST}:K{LAST}", '"' + ",".join(INTENTS) + '"', strict=False,
                    prompt="Optional: paste the label from an intent classifier. Overrides the rule-based intent.")
    number_validation(km, f"C{FIRST}:C{LAST}", 0)
    number_validation(km, f"D{FIRST}:D{LAST}", 0, 100)
    number_validation(km, f"F{FIRST}:F{LAST}", 0)
    band(km, f"A{FIRST}:L{LAST}", "A", FIRST)
    flag_rule(km, f"R{FIRST}:R{LAST}", f'$R{FIRST}="Yes"', F_RED, bold=True)
    for col in ("T", "U", "W", "X"):
        flag_rule(km, f"{col}{FIRST}:{col}{LAST}", f'${col}{FIRST}="Yes"', F_HI)
    flag_rule(km, f"Z{FIRST}:Z{LAST}", f'AND($Z{FIRST}<>"",$Z{FIRST}<>"OK",$Z{FIRST}<>"Not mapped")', F_HI, bold=True)
    flag_rule(km, f"Z{FIRST}:Z{LAST}", f'$Z{FIRST}="Not mapped"', F_LAV)
    widths(km, {"A": 38, "B": 18, "C": 11, "D": 11, "E": 40, "F": 10, "G": 44, "H": 14, "I": 22, "J": 11, "K": 15, "L": 30,
                "M": 30, "N": 15, "O": 14, "P": 10, "Q": 11, "R": 14, "S": 12, "T": 12, "U": 12, "V": 12, "W": 11,
                "X": 13, "Y": 8, "Z": 46, "AA": 9, "AB": 9})
    km.freeze_panes = "B3"
    km.auto_filter.ref = f"A2:AB{LAST}"

    # ============================================================ URL Summary
    title(us, "A1", "URL Summary: one row per target URL")
    hdr = ["#", "Target URL", "Page type", "Page status", "Keywords mapped", "Primary keywords", "Search volume (all mapped)",
           "Avg priority", "Cannibalised keywords", "Intent mismatches", "Status"]
    header_row(us, 2, hdr, 1, "calc")
    g, i_, h_, j_ = (f"{KM}!$G${FIRST}:$G${LAST}", f"{KM}!$I${FIRST}:$I${LAST}", f"{KM}!$H${FIRST}:$H${LAST}", f"{KM}!$J${FIRST}:$J${LAST}")
    aa = f"{KM}!$AA${FIRST}:$AA${LAST}"
    for r in range(3, URL_LAST + 1):
        e = esc(f"B{r}")
        us[f"A{r}"] = f"=ROW()-2"
        us[f"B{r}"] = f'=IFERROR(INDEX({g},MATCH(A{r},{aa},0)),"")'
        us[f"C{r}"] = f'=IF(B{r}="","",INDEX({i_},MATCH(A{r},{aa},0))&"")'
        us[f"D{r}"] = f'=IF(B{r}="","",INDEX({h_},MATCH(A{r},{aa},0))&"")'
        us[f"E{r}"] = f'=IF(B{r}="","",COUNTIF({g},{e}))'
        us[f"F{r}"] = f'=IF(B{r}="","",COUNTIFS({g},{e},{j_},"Primary"))'
        us[f"G{r}"] = f'=IF(B{r}="","",SUMIFS({KM}!$C${FIRST}:$C${LAST},{g},{e}))'
        us[f"H{r}"] = f'=IF(B{r}="","",IFERROR(ROUND(AVERAGEIFS({KM}!$P${FIRST}:$P${LAST},{g},{e}),0),""))'
        us[f"I{r}"] = f'=IF(B{r}="","",COUNTIFS({g},{e},{KM}!$R${FIRST}:$R${LAST},"Yes"))'
        us[f"J{r}"] = f'=IF(B{r}="","",COUNTIFS({g},{e},{KM}!$W${FIRST}:$W${LAST},"Yes"))'
        us[f"K{r}"] = (f'=IF(B{r}="","",IF(F{r}=0,"No primary keyword",IF(F{r}>Settings!$B$4,"Too many primary keywords",'
                       f'IF(I{r}+J{r}>0,"Check flags","OK"))))')
    band(us, f"A3:K{URL_LAST}", "B", 3)
    flag_rule(us, f"K3:K{URL_LAST}", '$K3="Too many primary keywords"', F_RED, bold=True)
    flag_rule(us, f"K3:K{URL_LAST}", 'OR($K3="Check flags",$K3="No primary keyword")', F_HI)
    widths(us, {"A": 6, "B": 52, "C": 22, "D": 14, "E": 11, "F": 11, "G": 14, "H": 10, "I": 13, "J": 12, "K": 26})
    us.freeze_panes = "C3"
    us.auto_filter.ref = f"A2:K{URL_LAST}"

    # ============================================================ Cluster Summary
    title(cs, "A1", "Cluster Summary")
    header_row(cs, 2, ["#", "Cluster", "Keywords", "Search volume", "Mapped to a URL", "Not mapped", "Cannibalised keywords",
                       "Intent mismatches", "Avg priority"], 1, "calc")
    b_ = f"{KM}!$B${FIRST}:$B${LAST}"
    ab = f"{KM}!$AB${FIRST}:$AB${LAST}"
    for r in range(3, CL_LAST + 1):
        e = esc(f"B{r}")
        cs[f"A{r}"] = "=ROW()-2"
        cs[f"B{r}"] = f'=IFERROR(INDEX({b_},MATCH(A{r},{ab},0)),"")'
        cs[f"C{r}"] = f'=IF(B{r}="","",COUNTIF({b_},{e}))'
        cs[f"D{r}"] = f'=IF(B{r}="","",SUMIFS({KM}!$C${FIRST}:$C${LAST},{b_},{e}))'
        cs[f"E{r}"] = f'=IF(B{r}="","",SUMPRODUCT(({b_}=B{r})*(LEN({g})>0)))'
        cs[f"F{r}"] = f'=IF(B{r}="","",C{r}-E{r})'
        cs[f"G{r}"] = f'=IF(B{r}="","",COUNTIFS({b_},{e},{KM}!$R${FIRST}:$R${LAST},"Yes"))'
        cs[f"H{r}"] = f'=IF(B{r}="","",COUNTIFS({b_},{e},{KM}!$W${FIRST}:$W${LAST},"Yes"))'
        cs[f"I{r}"] = f'=IF(B{r}="","",IFERROR(ROUND(AVERAGEIFS({KM}!$P${FIRST}:$P${LAST},{b_},{e}),0),""))'
    band(cs, f"A3:I{CL_LAST}", "B", 3)
    flag_rule(cs, f"F3:H{CL_LAST}", 'AND(ISNUMBER(F3),F3>0)', F_HI)
    widths(cs, {"A": 6, "B": 30, "C": 11, "D": 14, "E": 14, "F": 12, "G": 14, "H": 13, "I": 11})
    cs.freeze_panes = "C3"
    cs.auto_filter.ref = f"A2:I{CL_LAST}"

    # ============================================================ Dashboard
    title(dash, "A1", "Keyword map dashboard")
    note(dash, "A2", "Counts update as you edit Keyword Map. Yellow numbers above zero are worth a look.")
    a = f"{KM}!$A${FIRST}:$A${LAST}"
    kpis = [
        ("Keywords", f'=SUMPRODUCT((LEN({a})>0)*1)'),
        ("Mapped to a URL", f'=SUMPRODUCT((LEN({a})>0)*(LEN({g})>0))'),
        ("Not mapped yet", "=B4-B5"),
        ("Target URLs", f"=MAX({aa})"),
        ("Planned (new) URLs", "=COUNTIF('URL Summary'!$D$3:$D$%d,\"Planned\")" % URL_LAST),
        ("Clusters", f"=MAX({ab})"),
        ("Cannibalised keyword rows", f'=COUNTIF({KM}!$R${FIRST}:$R${LAST},"Yes")'),
        ("URLs with too many primary keywords", "=COUNTIF('URL Summary'!$K$3:$K$%d,\"Too many primary keywords\")" % URL_LAST),
        ("URLs with no primary keyword", "=COUNTIF('URL Summary'!$K$3:$K$%d,\"No primary keyword\")" % URL_LAST),
        ("URLs spanning 2+ clusters (keyword rows)", f'=COUNTIF({KM}!$U${FIRST}:$U${LAST},"Yes")'),
        ("Intent mismatches", f'=COUNTIF({KM}!$W${FIRST}:$W${LAST},"Yes")'),
        ("Ranking URL is not the target", f'=COUNTIF({KM}!$X${FIRST}:$X${LAST},"Yes")'),
        ("Unclassified intent", f'=COUNTIF({KM}!$O${FIRST}:$O${LAST},"Unclassified")'),
    ]
    header_row(dash, 3, ["Metric", "Count"], 1, "calc")
    for i, (k, v) in enumerate(kpis, start=4):
        dash.cell(row=i, column=1, value=k).font = font(10)
        c = dash.cell(row=i, column=2, value=v); c.font = font(12, True, DARK); c.alignment = Alignment(horizontal="center")
    flag_rule(dash, "B10:B16", "AND(ISNUMBER(B10),B10>0)", F_HI)
    header_row(dash, 3, ["Final intent", "Keywords", "Search volume"], 4, "calc")
    for i, it in enumerate(INTENTS + ["Unclassified"], start=4):
        dash.cell(row=i, column=4, value=it).font = font(10)
        dash.cell(row=i, column=5, value=f'=COUNTIF({KM}!$O${FIRST}:$O${LAST},D{i})')
        dash.cell(row=i, column=6, value=f'=SUMIFS({KM}!$C${FIRST}:$C${LAST},{KM}!$O${FIRST}:$O${LAST},D{i})')
    header_row(dash, 11, ["Page type", "Keywords", "Intent mismatches"], 4, "calc")
    for i in range(12):
        r = 12 + i
        dash.cell(row=r, column=4, value=f'=IF(Settings!D{5 + i}="","",Settings!D{5 + i})')
        dash.cell(row=r, column=5, value=f'=IF(D{r}="","",COUNTIF({i_},D{r}))')
        dash.cell(row=r, column=6, value=f'=IF(D{r}="","",COUNTIFS({i_},D{r},{KM}!$W${FIRST}:$W${LAST},"Yes"))')
    widths(dash, {"A": 40, "B": 12, "C": 4, "D": 26, "E": 12, "F": 16})

    # ============================================================ Start here
    s = start
    title(s, "A1", "Keyword Mapping Template")
    s["A2"] = "Keyword → URL with intent, priority and cannibalisation flags · by Lazarina Stoy · MLforSEO"
    s["A2"].font = font(11, False, "8448EF", True)
    r = lines(s, 4, [
        "What it does",
        "Maps every keyword (or cluster) to one existing or planned URL, labels its search intent, scores priority, and flags",
        "the problems that cause cannibalisation: a keyword mapped to two pages, a page with too many primary keywords, a page",
        "spanning several clusters, a keyword whose intent doesn't fit the page type, and a ranking URL that isn't the one you chose.",
        "It is the step between keyword clustering and content briefs.",
    ], bold_first=True)
    r = lines(s, r + 1, [
        "How to use it",
        "1. Settings: replace the Navigational words with your brand and product names. Check the page types and the intent-fit grid.",
        "2. Keyword Map: delete the sample rows (columns A-L only), then paste your clustered keywords: keyword, cluster, volume, KD.",
        "   Optional: current ranking URL and position from Search Console or your rank tracker, and an intent label from a classifier.",
        "3. For each keyword pick a Target URL, Page status, Page type and Keyword role (one Primary per URL, the rest Secondary).",
        "4. Read the Flags column. Fix cannibalisation first (choose one URL per keyword), then overloaded URLs, then intent mismatches.",
        "5. URL Summary and Cluster Summary list every target URL and cluster with counts. Dashboard shows the totals.",
        "6. Planned URLs with a Primary keyword are your content briefs: take them into the MLforSEO content brief template.",
    ], bold_first=True)
    r = lines(s, r + 1, [
        "How the flags work",
        "Cannibalisation: the same keyword (case and spaces ignored) appears on another row with a different target URL.",
        "Too many primaries on URL: the target URL has more Primary keywords than the limit on Settings (default 1).",
        "URL spans 2+ clusters: the same target URL also holds keywords from another cluster.",
        "Intent mismatch: the keyword's final intent is marked 'No' for that page type in the Settings grid.",
        "Ranking URL is not the target: Search Console or your tracker shows a different page ranking (trailing slash and case ignored).",
        "   On a Planned page this is expected: it tells you which existing page will compete until you relink or redirect it.",
        "Intent: a pasted label wins; otherwise whole-word rules from Settings (Navigational > Transactional > Commercial > Informational).",
        "Priority score (0-100): volume points (log scale) + ease points (100 - KD) + intent points. All weights live on Settings.",
    ], bold_first=True)
    r = legend(s, r + 1)
    r = lines(s, r, [
        "Good to know",
        "Works in Google Sheets and Excel: formulas are INDEX/MATCH, COUNTIFS, SUMIFS, SUMPRODUCT and IF. Room for 1,000 keywords",
        "(copy the last formula row down for more). Rule-based intent is a first pass, not ground truth. COUNTIFS ignores case.",
    ], bold_first=True)
    r = credits(s, r + 1, "Read: How to turn semantic keyword research into an actionable SEO workflow (MLforSEO)", POST_URL,
                "Course: Semantic ML-enabled Keyword Research (MLforSEO Academy)", COURSE_URL,
                [("Free intent classifier (code): github.com/lazarinastoy/mlforseo-code/search-intent-regex-classifier",
                  "https://github.com/lazarinastoy/mlforseo-code/tree/main/search-intent-regex-classifier")])
    widths(s, {"A": 30, "B": 110})
    s.sheet_view.showGridLines = False
    for ws in (dash,):
        ws.sheet_view.showGridLines = False
    s.sheet_properties.tabColor = DARK
    km.sheet_properties.tabColor = "8448EF"
    dash.sheet_properties.tabColor = "FFDE59"
    wb.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    build()
